"""pptx, pdf, txt를 출처가 남는 청크로 읽는다."""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from insight_distiller.models import Chunk
from insight_distiller.textutil import normalize_keep_lines, repair_ocr

QUALITY = {
    "pptx": 0.95,
    "txt": 0.90,
    "pdf_text": 0.78,
    "pdf_ocr": 0.60,
    "transcript": 0.70,
}


def ingest_dir(path: Path, ocr: bool = True) -> list[Chunk]:
    files = sorted(item for item in path.iterdir() if item.is_file() and not item.name.startswith("."))
    chunks: list[Chunk] = []
    for file in files:
        suffix = file.suffix.lower()
        if suffix == ".pptx":
            chunks.extend(_pptx(file))
        elif suffix == ".pdf":
            chunks.extend(_pdf(file, ocr=ocr))
        elif suffix in {".txt", ".md"}:
            chunks.extend(_txt(file))
        else:
            continue
    return chunks


def document_stats(chunks: list[Chunk]) -> list[dict]:
    stats: dict[str, dict] = {}
    for chunk in chunks:
        row = stats.setdefault(
            chunk.source,
            {"source": chunk.source, "chars": 0, "chunks": 0, "kinds": []},
        )
        row["chars"] += chunk.chars
        row["chunks"] += 1
        if chunk.kind not in row["kinds"]:
            row["kinds"].append(chunk.kind)
    return list(stats.values())


def _pptx(path: Path) -> list[Chunk]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    chunks: list[Chunk] = []
    for index, slide in enumerate(deck.slides, start=1):
        lines: list[str] = []
        for shape in _shapes(slide.shapes, MSO_SHAPE_TYPE):
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        lines.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        text = normalize_keep_lines("\n".join(lines))
        if text:
            chunks.append(Chunk(path.name, "pptx", f"slide {index}", text, QUALITY["pptx"]))
    return chunks


def _shapes(shapes, group_type):
    for shape in shapes:
        if shape.shape_type == group_type.GROUP:
            yield from _shapes(shape.shapes, group_type)
        else:
            yield shape


def _pdf(path: Path, ocr: bool) -> list[Chunk]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks: list[Chunk] = []
    for index, page in enumerate(reader.pages, start=1):
        text = normalize_keep_lines(page.extract_text() or "")
        if len(text) >= 20:
            chunks.append(Chunk(path.name, "pdf_text", f"page {index}", text, QUALITY["pdf_text"]))
    if ocr and _ocr_available():
        chunks.extend(_ocr_pdf(path))
    return chunks


def _ocr_available() -> bool:
    return shutil.which("pdftoppm") is not None and shutil.which("tesseract") is not None


def _ocr_pdf(path: Path) -> list[Chunk]:
    digest = hashlib.sha256(path.read_bytes() + b"ocr-v3").hexdigest()[:16]
    cache = Path(tempfile.gettempdir()) / "insight-ocr-cache" / digest
    cache.mkdir(parents=True, exist_ok=True)
    chunks: list[Chunk] = []
    with tempfile.TemporaryDirectory(prefix="insight-pdf-") as directory:
        prefix = str(Path(directory) / "page")
        subprocess.run(
            ["pdftoppm", "-png", "-r", "200", str(path), prefix],
            check=True,
            capture_output=True,
        )
        images = sorted(Path(directory).glob("page*.png"))
        for index, image in enumerate(images, start=1):
            cached = cache / f"{index}.txt"
            if cached.exists():
                text = cached.read_text(encoding="utf-8")
            else:
                completed = subprocess.run(
                    ["tesseract", str(image), "stdout", "-l", "kor+eng", "--psm", "6"],
                    check=True,
                    capture_output=True,
                    text=True,
                )
                text = repair_ocr(normalize_keep_lines(completed.stdout))
                cached.write_text(text, encoding="utf-8")
            if len(text) >= 20:
                chunks.append(Chunk(path.name, "pdf_ocr", f"page {index}", text, QUALITY["pdf_ocr"]))
    return chunks


def _txt(path: Path) -> list[Chunk]:
    raw = normalize_keep_lines(path.read_text(encoding="utf-8"))
    parts = re.split(r"Quick 코멘트\s*", raw)
    chunks: list[Chunk] = []
    seen: set[str] = set()
    for index, part in enumerate(parts):
        text = part.strip()
        if len(text) < 40:
            continue
        fingerprint = re.sub(r"\s+", "", text)[:400]
        if fingerprint in seen:
            continue
        seen.add(fingerprint)
        transcript = _is_transcript(text)
        kind = "transcript" if transcript else "txt"
        quality = QUALITY["transcript"] if transcript else QUALITY["txt"]
        locator = "transcript" if transcript else f"comment {index}"
        chunks.append(Chunk(path.name, kind, locator, text, quality))
    if not chunks:
        chunks.append(Chunk(path.name, "txt", "body", raw, QUALITY["txt"]))
    return chunks


def _is_transcript(text: str) -> bool:
    clock = len(re.findall(r"(?m)^\d+:\d+$", text))
    seconds = len(re.findall(r"(?m)^\d+초$", text))
    return len(text) > 15000 and (clock + seconds) > 8
