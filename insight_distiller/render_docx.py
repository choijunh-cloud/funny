"""Lecture-note docx for a distilled report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor

from insight_distiller.distill import Report

KR_FONT = "맑은 고딕"
NAVY = RGBColor(0x0F, 0x20, 0x43)
NAVY2 = RGBColor(0x1E, 0x40, 0x7C)
GOLD = RGBColor(0xB8, 0x94, 0x3A)
GRAY = RGBColor(0x4B, 0x55, 0x63)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX = "0F2043"
NAVY2_HEX = "1E407C"
GOLD_HEX = "B8943A"
LIGHT_HEX = "EEF2F8"
ROW_HEX = "F7F9FC"
WHITE_HEX = "FFFFFF"


def _font(run, size=11, bold=False, color=DARK):
    run.font.name = KR_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), KR_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), KR_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def _shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _margins(cell, top=60, bottom=60, left=80, right=80):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f"</w:tcMar>"
        )
    )


def _borders(table, color="D5DCE6"):
    tbl_pr = table._tbl.tblPr
    tbl_pr.append(
        parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f"</w:tblBorders>"
        )
    )


def _cell(cell, text, size=9.5, bold=False, color=DARK, align="left", fill=None):
    if fill:
        _shade(cell, fill)
    cell.text = ""
    align_enum = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }[align]
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.alignment = align_enum
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.line_spacing = 1.08
        run = para.add_run(line)
        _font(run, size=size, bold=bold, color=color)
    _margins(cell)


class Notes:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Mm(210)
        sec.page_height = Mm(297)
        sec.left_margin = Mm(16)
        sec.right_margin = Mm(16)
        sec.top_margin = Mm(16)
        sec.bottom_margin = Mm(16)
        normal = self.doc.styles["Normal"]
        normal.font.name = KR_FONT
        normal.font.size = Pt(11)
        normal.font.color.rgb = DARK
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
        header = sec.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("투자 인사이트 증류  ·  퀵코멘트 · 녹취")
        _font(run, size=8.5, color=GRAY)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("원문 추출  ·  숫자는 코멘트 표기  ·  매수·매도 권유 아님  ·  ")
        _font(run, size=8, color=GRAY)
        footer._p.append(
            parse_xml(
                f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE ">'
                f'<w:r><w:rPr><w:sz w:val="16"/><w:color w:val="4B5563"/>'
                f'<w:rFonts w:ascii="{KR_FONT}" w:hAnsi="{KR_FONT}" w:eastAsia="{KR_FONT}"/>'
                f"</w:rPr><w:t></w:t></w:r></w:fldSimple>"
            )
        )
        core = self.doc.core_properties
        core.title = "투자 인사이트 증류"
        core.author = "준혁"
        core.subject = "퀵코멘트와 방송 녹취에서 추출한 투자 인사이트"

    def p(self, text, size=11, bold=False, color=DARK, align="left", before=0, after=6):
        para = self.doc.add_paragraph()
        para.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
        }[align]
        para.paragraph_format.space_before = Pt(before)
        para.paragraph_format.space_after = Pt(after)
        para.paragraph_format.line_spacing = 1.15
        run = para.add_run(text)
        _font(run, size=size, bold=bold, color=color)
        return para

    def h1(self, text):
        para = self.p(text, size=16, bold=True, color=NAVY, before=12, after=6)
        para._p.get_or_add_pPr().append(
            parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="12" w:space="3" w:color="{NAVY_HEX}"/>'
                f"</w:pBdr>"
            )
        )

    def h2(self, text):
        self.p(text, size=13, bold=True, color=NAVY2, before=8, after=3)

    def bullet(self, text, lead=None):
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.55)
        para.paragraph_format.first_line_indent = Cm(-0.35)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing = 1.12
        run = para.add_run("• ")
        _font(run, size=10.5, color=NAVY2)
        if lead:
            run = para.add_run(lead)
            _font(run, size=10.5, bold=True, color=DARK)
        run = para.add_run(text)
        _font(run, size=10.5, color=DARK)

    def callout(self, title, body):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        _shade(cell, LIGHT_HEX)
        cell._tc.get_or_add_tcPr().append(
            parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="nil"/>'
                f'<w:left w:val="single" w:sz="28" w:space="0" w:color="{GOLD_HEX}"/>'
                f'<w:bottom w:val="nil"/>'
                f'<w:right w:val="nil"/>'
                f"</w:tcBorders>"
            )
        )
        _margins(cell, top=80, bottom=80, left=120, right=120)
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        run = p1.add_run(title)
        _font(run, size=10, bold=True, color=NAVY)
        for line in body:
            para = cell.add_paragraph()
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.line_spacing = 1.12
            run = para.add_run(line)
            _font(run, size=10.5, color=DARK)
        self.p("", size=6, after=4)

    def table(self, headers, rows, widths):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(table)
        for i, header in enumerate(headers):
            _cell(table.rows[0].cells[i], header, size=9, bold=True, color=WHITE, align="center", fill=NAVY_HEX)
        for r_i, row in enumerate(rows):
            fill = ROW_HEX if r_i % 2 else WHITE_HEX
            for c_i, val in enumerate(row):
                _cell(
                    table.rows[r_i + 1].cells[c_i],
                    val,
                    size=8.5,
                    bold=c_i == 0,
                    align="left" if c_i == 0 else "center",
                    fill=fill,
                )
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
        self.p("", size=4, after=2)


def render_docx(report: Report, path: Path) -> None:
    notes = Notes()
    stats = report.stats
    notes.p("퀵코멘트 · 방송 녹취 증류", size=11, bold=True, color=GOLD, align="center", after=2)
    notes.p("투자 인사이트", size=22, bold=True, color=NAVY, align="center", after=2)
    notes.p(
        f"퀵코멘트 {stats['quick_kept']}개 · 대담 {len(report.dialogs)}편 · 한 노트 {stats.get('unified_bullets', 0)}문장",
        size=10.5,
        color=GRAY,
        align="center",
        after=8,
    )
    notes.callout(
        "한 줄",
        [report.headline or "원문에서 프레임 문장을 찾지 못했다."],
    )
    if report.sources:
        notes.p("출처 표기  " + " · ".join(report.sources), size=10, color=GRAY, after=4)
    notes.p(
        "퀵코멘트와 방송 대담을 한 노트로 합쳤다. 같은 논지는 한 번만 남기고 출처를 붙인다.",
        size=10,
        color=GRAY,
        after=6,
    )

    for index, section in enumerate(report.sections, start=1):
        notes.h1(f"{index}. {section.title}")
        if section.lead:
            notes.p(section.lead, size=11, after=6)
        if section.table:
            notes.table(
                ["종목", "가격", "앞 구간", "뒤 구간"],
                [
                    [
                        row.get("name", ""),
                        row.get("price", ""),
                        row.get("y26", "") or "—",
                        row.get("y27", "") or "—",
                    ]
                    for row in section.table
                ],
                [3.2, 3.6, 5.5, 5.5],
            )
        for bullet in section.bullets:
            if bullet.time:
                lead = f"{bullet.time}  "
            elif bullet.source == "transcript":
                lead = "녹취  "
            elif bullet.source == "extract":
                lead = "추출  "
            else:
                lead = ""
            notes.bullet(bullet.text, lead=lead or None)

    if report.checklist:
        notes.h1("확인할 조건")
        for item in report.checklist:
            notes.bullet(item)

    path.parent.mkdir(parents=True, exist_ok=True)
    notes.doc.save(str(path))
