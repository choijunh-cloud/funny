"""방송 전사에서 오늘의 논리와 투자 판단을 증류한다.

STT 오탈자를 바로잡지 않는다. 화자가 결론으로 미는 구절을 테마로 모으고,
그 갈림을 표로 옮긴다. 숫자는 맞추지 않는다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from insights.intent import factiness

TIMESTAMP = re.compile(r"\d+\s*시간\s*\d+\s*분(?:\s*\d+\s*초)?")
ENDING = re.compile(r"(니다|거든요|잖아요|겁니다|습니다|데요|거죠|거예요|니까요|아니죠|[.!?])\s+")
LEAD = re.compile(r"^(?:그래서|그니까|그러니까|근데|그러나|일단은|일단|자)\s+")
NOISE = (
    "카카오톡",
    "투자 판단의 책임",
    "감사합니다",
    "감사하고",
    "안녕하",
    "어서 오십",
    "잠시 후",
    "다음 장 보",
    "정리해 볼까요",
    "머니랩",
    "보고서 형식",
    "시각화 도표",
    "투자자분들께",
)
SETUP = (
    "보겠습니다",
    "볼까요",
    "주시죠",
    "알아보겠",
    "체크해 두",
    "살펴보",
    "확인해 보면",
    "전해 주시",
)
MARKERS = (
    ("결국", 5),
    ("혁명", 4),
    ("핵심", 4),
    ("아니라", 5),
    ("아니죠", 5),
    ("아니다", 3),
    ("봐야", 3),
    ("보셔야", 3),
    ("때문에", 3),
    ("문제가", 3),
    ("경고", 3),
    ("손절", 4),
    ("버티", 3),
    ("중요", 2),
    ("변수", 2),
    ("증설", 2),
    ("수주", 2),
)


@dataclass(frozen=True)
class ThemeSpec:
    key: str
    label: str
    primary: tuple[str, ...]
    secondary: tuple[str, ...]


THEMES: tuple[ThemeSpec, ...] = (
    ThemeSpec(
        "rates",
        "금리보다 그 위의 성장",
        ("금리 이상", "금리보다", "혁명", "역전", "윌리엄스", "윌리어스", "서두를", "한 번만"),
        ("금리", "인상", "10년", "성장"),
    ),
    ThemeSpec(
        "oil",
        "유가가 풀리면 금리와 주식이 따라온다",
        ("유가 안정", "유가가 안정"),
        ("유가", "WTI", "원유"),
    ),
    ThemeSpec(
        "korea",
        "외국인은 팔고 장비와 대형주가 버틴다",
        ("외국인", "배당락", "6,500", "6500", "갇혀", "조단위", "장비 관련"),
        ("코스피", "코스닥", "소부장"),
    ),
    ThemeSpec(
        "ai",
        "빨리 가자는 경쟁과 늦추자는 경고가 같이 있다",
        ("속도를 조절", "속도 조절", "더 빨리", "경쟁과", "에이전트", "다츠", "닷츠", "다치", "다스"),
        ("뮤즈", "앤트로픽", "엔트로픽", "안전 문제", "인공지능"),
    ),
    ThemeSpec(
        "memory",
        "메모리는 아직 사이클, 차가운 가격은 수급",
        ("천민", "사이클 산업", "수급상", "못 팔", "없어서 못"),
        ("마이크론", "마이크로", "쇼티지", "가이던스", "사이클"),
    ),
    ThemeSpec(
        "hbm",
        "잘해도 업황, 못해도 몫을 나눈다",
        ("나눠 먹", "반대 급", "키가 HBM", "HBM으로", "수율"),
        ("HBM", "점유", "블랙웰", "베라"),
    ),
    ThemeSpec(
        "valuation",
        "전통 사이클 배수와 장기 계약 배수",
        ("뉴노멀", "PBR", "장기 계약", "B2B"),
        ("PER", "전통", "사이클"),
    ),
    ThemeSpec(
        "bonds",
        "채권을 찍는 동안은 경고가 아니다",
        ("채권 발행", "사채", "사체", "다컴", "닷컴", "CDS", "정상"),
        ("채권", "하이퍼"),
    ),
    ThemeSpec(
        "materials",
        "장비는 앞서고 소재는 늦다",
        ("장비주", "가동률", "그린필드", "웨이퍼", "소재 로테"),
        ("소재", "기판"),
    ),
    ThemeSpec(
        "action",
        "박스에선 주도주를 버틴다",
        ("수면제", "손절", "매물대", "엇박", "190만"),
        ("주도주", "버티", "신규", "현금"),
    ),
    ThemeSpec(
        "names",
        "개별 호재는 증설과 수주가 보일 때",
        ("삼성전기", "한미반도체", "한미 반도체", "HLB", "빅웨이브"),
        ("증설", "수주", "기판", "허가"),
    ),
)


@dataclass
class StudyPiece:
    theme: str
    label: str
    point: str
    support: list[str]
    order: int
    insight: str = ""
    notes: list[str] = field(default_factory=list)


@dataclass
class Study:
    title: str
    intro: str
    logic: list[str]
    pieces: list[StudyPiece]
    chain: list[tuple[str, str, str]] = field(default_factory=list)
    forks: list[tuple[str, str, str]] = field(default_factory=list)
    actions: list[tuple[str, str]] = field(default_factory=list)
    stats: dict[str, int] = field(default_factory=dict)


def _atoms(text: str) -> list[str]:
    cleaned = TIMESTAMP.sub(" ", text)
    cleaned = re.sub(r"\s+", " ", cleaned)
    parts = ENDING.split(cleaned)
    merged: list[str] = []
    if len(parts) == 1:
        merged = parts
    else:
        for index in range(0, len(parts) - 1, 2):
            merged.append(f"{parts[index]}{parts[index + 1]}")
        if len(parts) % 2 == 1 and parts[-1].strip():
            merged.append(parts[-1])
    out: list[str] = []
    for part in merged:
        sentence = re.sub(r"\s+", " ", part).strip(" .")
        chunks = [sentence]
        if len(sentence) > 220:
            chunks = [chunk.strip() for chunk in re.split(r"[,，]\s*", sentence)]
        for chunk in chunks:
            if len(chunk) < 12 or len(chunk) > 220:
                continue
            if any(noise in chunk for noise in NOISE):
                continue
            out.append(chunk)
    return out


def _spans(atoms: list[str]) -> list[str]:
    spans: list[str] = []
    for start, atom in enumerate(atoms):
        joined = ""
        for follow in atoms[start : start + 3]:
            joined = f"{joined} {follow}".strip()
            if len(joined) > 190:
                break
            if len(joined) >= 24:
                spans.append(joined)
    return spans


def _hits(sentence: str, words: tuple[str, ...]) -> int:
    return sum(1 for word in words if word in sentence)


def _theme_of(sentence: str) -> ThemeSpec | None:
    ranked = []
    for spec in THEMES:
        primary = _hits(sentence, spec.primary)
        if primary:
            ranked.append((primary, _hits(sentence, spec.secondary), spec))
    if not ranked:
        secondary = []
        for spec in THEMES:
            hits = _hits(sentence, spec.secondary)
            if hits:
                secondary.append((hits, spec))
        if not secondary:
            return None
        secondary.sort(key=lambda item: item[0], reverse=True)
        if len(secondary) > 1 and secondary[0][0] == secondary[1][0]:
            return None
        return secondary[0][1]
    ranked.sort(key=lambda item: (item[0], item[1]), reverse=True)
    if len(ranked) > 1 and ranked[0][0] < ranked[1][0] * 2:
        return None
    return ranked[0][2]


def _score(sentence: str, spec: ThemeSpec | None) -> int:
    value = 0
    for marker, points in MARKERS:
        if marker in sentence:
            value += points
    if spec is not None:
        value += 6 * _hits(sentence, spec.primary)
    filler = len(re.findall(r"어 |그니까|뭐 |자,|음 ", sentence))
    value -= min(8, filler * 2)
    if factiness(sentence) >= 0.08:
        value -= 4
    if any(tail in sentence for tail in SETUP):
        value -= 4
    if sentence.count("?") >= 2:
        value -= 3
    if "거 같" in sentence or "것 같" in sentence:
        value -= 4
    if "공시" in sentence or "생산 능력" in sentence:
        value += 3
    if "상승 마감" in sentence or "급등" in sentence:
        value -= 3
    length = len(sentence)
    if 36 <= length <= 140:
        value += 2
    elif length > 170:
        value -= 2
    return value


def _point(sentence: str, spec: ThemeSpec) -> str:
    text = LEAD.sub("", re.sub(r"\s+", " ", sentence)).strip(" .")
    primaries = [word for word in spec.primary if word in text]
    anchors = primaries or [word for word in spec.secondary if word in text]
    if len(text) <= 110 or not anchors:
        return text[:130].rstrip(" ,")
    index = min(text.find(anchor) for anchor in anchors)
    start = max(0, index - 18)
    snap = text.rfind(" ", 0, start + 1)
    if snap >= 0 and index - snap <= 20:
        start = snap + 1
    end = min(len(text), index + 100)
    snippet = text[start:end].strip(" ,.")
    snippet = re.sub(r"^(?:니다|습니다|거든요|잖아요|데요)\s*", "", snippet)
    snippet = re.sub(r"^(?:이|은|는|을|를|가|과|와|요|고|서|면)\s+", "", snippet)
    for tail in ("체크해", "알아보겠", "내일 새벽"):
        if tail in snippet:
            snippet = snippet.split(tail)[0].strip(" ,.")
    if end < len(text):
        cut = snippet.rfind(" ")
        if cut >= 36:
            snippet = snippet[:cut]
    return snippet.strip(" ,.")


def _similar(left: str, right: str) -> bool:
    from difflib import SequenceMatcher

    compact_left = re.sub(r"\s+", "", left)
    compact_right = re.sub(r"\s+", "", right)
    if len(compact_left) >= 18 and len(compact_right) >= 18 and (
        compact_left in compact_right or compact_right in compact_left
    ):
        return True
    short, long = (compact_left, compact_right) if len(compact_left) <= len(compact_right) else (compact_right, compact_left)
    if len(short) >= 28 and short[:28] in long:
        return True
    return SequenceMatcher(None, compact_left, compact_right).ratio() >= 0.72


def _keep(ranked: list[tuple[int, str]], spec: ThemeSpec) -> list[str]:
    kept: list[str] = []
    covered: set[str] = set()

    def take(sentence: str, need_primary: bool) -> bool:
        point = _point(sentence, spec)
        if len(point) < 16 or any(_similar(point, old) for old in kept):
            return False
        fresh = {word for word in spec.primary if word in point and word not in covered}
        if need_primary and not fresh:
            return False
        if kept and not fresh and not any(word in point and word not in covered for word in spec.secondary):
            return False
        kept.append(point)
        covered.update(word for word in (*spec.primary, *spec.secondary) if word in point)
        return True

    for _score_value, sentence in ranked:
        if take(sentence, False):
            break
    for _score_value, sentence in ranked:
        if len(kept) >= 4:
            break
        take(sentence, True)
    return kept


def _clauses(text: str) -> list[str]:
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"(매 ){2,}", "", text)
    parts = re.split(r"[?]\s*|,\s+| 그리고 | 근데 | 그래서 | 왜냐면 | 그러나 | 는데 ", text)
    clauses: list[str] = []
    for part in parts:
        clause = part.strip(" .")
        clause = re.sub(r"^(?:보면|하면|해서|니까|자|요|고|어)\s+", "", clause)
        if len(clause) >= 14:
            clauses.append(clause)
    return clauses or [text.strip(" .")]


def _distill_line(texts: list[str], spec: ThemeSpec) -> str:
    anchors = (*spec.primary, *spec.secondary)
    best = ""
    best_score = -10**9
    for text in texts:
        for clause in _clauses(text):
            if not any(word in clause for word in anchors):
                continue
            score = _score(clause, spec)
            if 18 <= len(clause) <= 76:
                score += 5
            elif len(clause) > 120:
                score -= 4
            if score > best_score:
                best_score = score
                best = clause
    if not best:
        best = texts[0]
    best = re.sub(r"^(?:아니라|보면|하면|일단은|그럼|그런데)\s*", "", best).strip()
    found = [word for word in spec.primary if word in best] or [word for word in spec.secondary if word in best]
    if found and len(best) > 76:
        index = min(best.find(word) for word in found)
        start = max(0, index - 8)
        snap = best.rfind(" ", 0, start + 1)
        if snap >= 0:
            start = snap + 1
        best = best[start:]
    best = _clip(best, 76)
    return re.sub(r"^(?:(?:은|는|이|가|을|를|의|에|와|과|도|만|요|고)\s+)+", "", best).strip()


def _clip(text: str, limit: int = 92) -> str:
    text = re.sub(r"\s+", " ", text).strip(" .")
    if len(text) <= limit:
        return text
    cut = text.rfind(" ", 0, limit)
    return (text[:cut] if cut >= 24 else text[:limit]).strip(" ,")


def _around(text: str, words: tuple[str, ...]) -> str:
    index = min(text.find(word) for word in words if word in text)
    start = max(0, index - 10)
    snap = text.rfind(" ", 0, start + 1)
    if snap >= 0 and index - snap <= 14:
        start = snap + 1
    end = min(len(text), index + 52)
    cut = text.rfind(" ", index, end)
    if cut > index + 16:
        end = cut
    return text[start:end].strip(" ,.")


def _side(piece: StudyPiece, words: tuple[str, ...], avoid: str = "") -> str:
    for text in (piece.point, *piece.support):
        if not any(word in text for word in words):
            continue
        clip = _around(text, words)
        if avoid and _similar(clip, avoid):
            continue
        if clip:
            return clip
    return ""


def _chain(pieces: list[StudyPiece]) -> list[tuple[str, str, str]]:
    oil = next((piece for piece in pieces if piece.theme == "oil"), None)
    rows: list[tuple[str, str, str]] = []
    if oil is not None:
        blob = f"{oil.point} {' '.join(oil.support)}"
        if "유가" in blob and "금리" in blob and any(word in blob for word in ("주식", "반등")):
            rows.append(("유가", "안정되면 금리가 안정되고", "주식이 반등할 자리"))
    rates = next((piece for piece in pieces if piece.theme == "rates"), None)
    if rates is not None and "혁명" in f"{rates.point} {' '.join(rates.support)}":
        revolution = _side(rates, ("혁명",))
        if revolution:
            rows.append(("혁명기", "자금 조달 수요가 금리를 밀어 올린다", revolution))
    return rows


def _forks(pieces: list[StudyPiece]) -> list[tuple[str, str, str]]:
    found = {piece.theme: piece for piece in pieces}
    specs = (
        ("valuation", "밸류를 가르는 눈", ("전통", "PBR", "사이클"), ("B2B", "뉴노멀", "PER", "장기")),
        ("hbm", "마이크론 실적의 두 갈래", ("잘해도", "잘하면", "업항", "업황"), ("나눠", "반대", "못 하면", "삐걱")),
        ("memory", "차가운 가격을 읽는 눈", ("천민", "사이클"), ("못 팔", "수급")),
        ("bonds", "채권이 경고가 되는 때", ("정상", "염려할 건 아니"), ("안 될", "사체", "사채", "다컴", "닷컴")),
        ("rates", "금리 상승을 읽는 법", ("성장", "혁명", "오르잖아요"), ("레벨", "빠졌", "역전")),
        ("materials", "장비와 소재의 시차", ("장비", "설비", "그린필드"), ("소재", "가동률", "웨이퍼")),
    )
    rows: list[tuple[str, str, str]] = []
    for key, title, left_words, right_words in specs:
        piece = found.get(key)
        if piece is None:
            continue
        left = _side(piece, left_words)
        right = _side(piece, right_words, avoid=left)
        if left and right and not _similar(left, right):
            rows.append((title, left, right))
    return rows


def _actions(pieces: list[StudyPiece]) -> list[tuple[str, str]]:
    piece = next((item for item in pieces if item.theme == "action"), None)
    if piece is None:
        return []
    blob_parts = [piece.point, *piece.support]
    rows: list[tuple[str, str]] = []
    mapping = (
        ("박스에서 사고 손절할 때", ("손절",)),
        ("주도주를 들고 있을 때", ("수면제", "버티")),
        ("이미 가진 뒤 현금을 더 넣을 때", ("매물대", "190만")),
        ("처음 들어가는 자리", ("신규",)),
    )
    used: list[str] = []
    for label, words in mapping:
        quote = ""
        for text in blob_parts:
            if not any(word in text for word in words):
                continue
            clip = _around(text, words)
            if any(_similar(clip, old) for old in used):
                continue
            quote = clip
            used.append(clip)
            break
        if quote:
            rows.append((label, quote))
    return rows


def distill_transcript(text: str) -> Study:
    atoms = _atoms(text)
    buckets: dict[str, list[tuple[int, str]]] = {spec.key: [] for spec in THEMES}
    for span in _spans(atoms):
        spec = _theme_of(span)
        if spec is None:
            continue
        score = _score(span, spec)
        if score < 4:
            continue
        buckets[spec.key].append((score, span))

    pieces: list[StudyPiece] = []
    for order, spec in enumerate(THEMES):
        ranked = sorted(buckets[spec.key], key=lambda item: item[0], reverse=True)
        kept = _keep(ranked, spec)
        if not kept:
            continue
        insight = _distill_line([kept[0]], spec)
        notes = []
        for extra in kept[1:]:
            note = _distill_line([extra], spec)
            if note and not _similar(note, insight) and not any(_similar(note, old) for old in notes):
                notes.append(note)
            if len(notes) >= 2:
                break
        pieces.append(
            StudyPiece(
                theme=spec.key,
                label=spec.label,
                point=kept[0],
                support=kept[1:],
                order=order,
                insight=insight,
                notes=notes,
            )
        )

    chain = _chain(pieces)
    forks = _forks(pieces)
    actions = _actions(pieces)
    return Study(
        title="오늘 학습, 화자들이 미는 논리",
        intro=(
            "오늘 방송에서 화자가 결론으로 민 말만 모았다. "
            "받아쓰기 오탈자는 고치지 않았고, 숫자는 맞추지 않았다. "
            "표는 그 말의 순서와 갈림이다."
        ),
        logic=[f"{piece.label}: {piece.insight}" for piece in pieces],
        pieces=pieces,
        chain=chain,
        forks=forks,
        actions=actions,
        stats={"sentences": len(atoms), "pieces": len(pieces)},
    )


def distill_transcript_path(path: str | Path) -> Study:
    return distill_transcript(Path(path).read_text(encoding="utf-8"))


def render_study_markdown(study: Study) -> str:
    lines = [
        f"# {study.title}",
        "",
        study.intro,
        "",
        f"말조각 {study.stats.get('sentences', 0)}개에서 논리 {study.stats.get('pieces', 0)}줄기로 좁힌 뒤, 줄기마다 결론 구절 한 줄로 증류했다.",
        "",
        "## 증류",
        "",
    ]
    for index, piece in enumerate(study.pieces, start=1):
        lines.append(f"{index}. **{piece.label}.** {piece.insight}")
        for note in piece.notes:
            lines.append(f"   - {note}")
    lines.extend(
        [
            "",
            "## 논리의 뼈대",
            "",
            "| 순서 | 줄기 | 증류한 말 |",
            "| --- | --- | --- |",
        ]
    )
    for index, piece in enumerate(study.pieces, start=1):
        lines.append(f"| {index} | {piece.label} | {piece.insight} |")
    lines.append("")
    if study.chain:
        lines.extend(["## 이어 붙인 인과", "", "| 앞에서 | 화자의 연결 | 뒤에서 |", "| --- | --- | --- |"])
        for left, mid, right in study.chain:
            lines.append(f"| {left} | {mid} | {right} |")
        lines.append("")
    if study.forks:
        lines.extend(["## 갈라 놓은 말", "", "| 갈림 | 한쪽 | 다른 쪽 |", "| --- | --- | --- |"])
        for title, left, right in study.forks:
            lines.append(f"| {title} | {left} | {right} |")
        lines.append("")
    if study.actions:
        lines.extend(["## 박스에서 하라는 말", "", "| 자리 | 화자의 말 |", "| --- | --- |"])
        for label, quote in study.actions:
            lines.append(f"| {label} | {quote} |")
        lines.append("")
    lines.extend(["## 줄기", ""])
    for piece in study.pieces:
        lines.append(f"### {piece.label}")
        lines.append("")
        lines.append(piece.insight)
        lines.append("")
        lines.append(f"받아쓴 말. {piece.point}")
        lines.append("")
        for item in piece.support:
            lines.append(f"- {item}")
        if piece.support:
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def write_study(study: Study, out_dir: Path) -> dict[str, Path]:
    import json

    out_dir.mkdir(parents=True, exist_ok=True)
    markdown_path = out_dir / "오늘학습.md"
    json_path = out_dir / "오늘학습.json"
    markdown_path.write_text(render_study_markdown(study), encoding="utf-8")
    payload = {
        "title": study.title,
        "intro": study.intro,
        "principle": "speaker-meaning-over-fact-check",
        "logic": [
            {
                "label": piece.label,
                "insight": piece.insight,
                "notes": piece.notes,
                "point": piece.point,
                "support": piece.support,
            }
            for piece in study.pieces
        ],
        "chain": [{"before": left, "link": mid, "after": right} for left, mid, right in study.chain],
        "forks": [{"split": title, "one": left, "other": right} for title, left, right in study.forks],
        "actions": [{"seat": label, "say": quote} for label, quote in study.actions],
        "stats": study.stats,
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    docx_path = out_dir / "오늘학습.docx"
    _write_docx(study, docx_path)
    return {"md": markdown_path, "json": json_path, "docx": docx_path}


def _write_docx(study: Study, path: Path) -> None:
    from docx import Document

    from insights.render import (
        DARK,
        GRAY,
        NAVY,
        NAVY_HEX,
        WHITE,
        _callout,
        _cell,
        _paragraph,
        _setup,
        _shade,
    )

    document = Document()
    _setup(document, study.title)
    header = document.sections[0].header.paragraphs[0]
    if header.runs:
        header.runs[0].text = "오늘 학습  ·  화자의 논리"
    _paragraph(document, study.title, size=20, bold=True, color=NAVY, space_after=4)
    _paragraph(document, study.intro, size=11, space_after=8)
    _paragraph(document, "증류", size=14, bold=True, color=NAVY, space_before=4, space_after=4)
    for index, piece in enumerate(study.pieces, start=1):
        _paragraph(document, f"{index}.  {piece.label}", size=11, bold=True, color=NAVY, space_after=1)
        _paragraph(document, piece.insight, size=11, space_after=1)
        for note in piece.notes:
            _paragraph(document, f"·  {note}", size=10.5, color=DARK, space_after=1)
        _paragraph(document, "", size=6, space_after=2)
    _paragraph(document, "논리의 뼈대", size=14, bold=True, color=NAVY, space_before=8, space_after=4)
    _grid(
        document,
        ["순서", "줄기", "증류한 말"],
        [[str(index), piece.label, piece.insight] for index, piece in enumerate(study.pieces, start=1)],
        NAVY_HEX,
        WHITE,
        _shade,
        _cell,
    )
    if study.chain:
        _paragraph(document, "이어 붙인 인과", size=14, bold=True, color=NAVY, space_before=12, space_after=4)
        _grid(document, ["앞에서", "화자의 연결", "뒤에서"], [list(row) for row in study.chain], NAVY_HEX, WHITE, _shade, _cell)
    if study.forks:
        _paragraph(document, "갈라 놓은 말", size=14, bold=True, color=NAVY, space_before=12, space_after=4)
        _grid(document, ["갈림", "한쪽", "다른 쪽"], [list(row) for row in study.forks], NAVY_HEX, WHITE, _shade, _cell)
    if study.actions:
        _paragraph(document, "박스에서 하라는 말", size=14, bold=True, color=NAVY, space_before=12, space_after=4)
        _grid(document, ["자리", "화자의 말"], [list(row) for row in study.actions], NAVY_HEX, WHITE, _shade, _cell)
    for piece in study.pieces:
        _paragraph(document, piece.label, size=14, bold=True, color=NAVY, space_before=12, space_after=4)
        _callout(document, "증류", piece.insight)
        _paragraph(document, f"받아쓴 말  ·  {piece.point}", size=10, color=GRAY, space_after=2)
        for item in piece.support:
            _paragraph(document, f"·  {item}", size=10.5, color=DARK, space_after=2)
    _paragraph(
        document,
        "숫자는 화자가 꺼낸 재료다. 맞추거나 고치지 않았다.",
        size=9,
        color=GRAY,
        space_before=8,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _grid(document, headers, rows, navy, white, shade, cell) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], navy)
        cell(table.rows[0].cells[index], header, bold=True, color=white, size=9)
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            cell(row.cells[index], value, size=8.5)
