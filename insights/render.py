"""증류 결과를 마크다운, JSON, 강의노트형 docx로 옮긴다."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

from insights.distill import Briefing, Meaning

NAVY = RGBColor(0x0F, 0x20, 0x43)
GOLD = RGBColor(0xB8, 0x94, 0x3A)
GRAY = RGBColor(0x4B, 0x55, 0x63)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
KR_FONT = "맑은 고딕"
NAVY_HEX = "0F2043"
LIGHT_HEX = "EEF2F8"
AMBER_HEX = "FFF8E7"


def render_markdown(briefing: Briefing, title: str) -> str:
    stats = briefing.stats
    lines = [
        f"# {title}",
        "",
        briefing.intro,
        "",
        (
            f"코멘트 {stats.get('comments', 0)}개 가운데 "
            f"한줄 평결 {stats.get('verdicts', 0)}개, "
            f"안내 멘트 {stats.get('logistics', 0)}개를 빼고 "
            f"판단 {stats.get('meanings', 0)}개로 합쳤다."
        ),
        "",
        "## 화자가 밀어붙인 판단",
        "",
    ]
    for index, meaning in enumerate(briefing.meanings, start=1):
        lines.append(f"{index}. **{meaning.label}.** {_first(meaning.point)}")
    lines.append("")
    for meaning in briefing.meanings:
        lines.extend(_meaning_md(meaning))
    if briefing.agenda:
        lines.append("## 화자가 짚은 일정")
        lines.append("")
        for item in briefing.agenda:
            lines.append(f"- {item.what}")
        lines.append("")
    if briefing.basket:
        lines.append("## 화자가 적어 둔 관심 바구니")
        lines.append("")
        for name, names in briefing.basket:
            lines.append(f"- {name}: {names}")
        lines.append("")
    if briefing.dropped:
        lines.append("## 판단에서 뺀 안내·한줄")
        lines.append("")
        lines.append("본문 판단을 대신하지 않는 멘트만 남겼다.")
        lines.append("")
        for item in briefing.dropped:
            lines.append(f"- {item}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_json(briefing: Briefing, title: str) -> str:
    payload = {
        "title": title,
        "intro": briefing.intro,
        "principle": "speaker-meaning-over-fact-check",
        "stats": briefing.stats,
        "judgments": [
            {
                "theme": meaning.theme,
                "label": meaning.label,
                "point": meaning.point,
                "alongside": meaning.alongside,
                "against": meaning.against,
                "watch": meaning.watch,
                "grounds": meaning.grounds,
                "names": meaning.names,
                "sources": meaning.sources,
            }
            for meaning in briefing.meanings
        ],
        "agenda": [item.what for item in briefing.agenda],
        "basket": [{"group": name, "names": names} for name, names in briefing.basket],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2) + "\n"


def render_docx(briefing: Briefing, title: str, path: Path) -> None:
    document = Document()
    _setup(document, title)
    _paragraph(document, title, size=20, bold=True, color=NAVY, space_after=4)
    _paragraph(
        document,
        "화자의 뜻  ·  숫자는 주장의 재료  ·  시세 검증 아님",
        size=10,
        color=GOLD,
        space_after=8,
    )
    _paragraph(document, briefing.intro, size=11, space_after=8)
    _paragraph(document, "화자가 밀어붙인 판단", size=14, bold=True, color=NAVY, space_before=8, space_after=6)
    for index, meaning in enumerate(briefing.meanings, start=1):
        _paragraph(document, f"{index}.  {meaning.label}", size=11, bold=True, color=NAVY, space_after=1)
        _paragraph(document, _first(meaning.point), size=11, space_after=4)
    for meaning in briefing.meanings:
        _paragraph(document, meaning.label, size=14, bold=True, color=NAVY, space_before=12, space_after=4)
        _callout(document, "말하려는 것", meaning.point)
        if meaning.against:
            _paragraph(document, f"밀어낸 해석  ·  {meaning.against}", size=10, color=GRAY, space_after=3)
        if meaning.alongside:
            _paragraph(document, "같이 실린 판단", size=11, bold=True, color=NAVY, space_before=4, space_after=2)
            for item in meaning.alongside:
                _paragraph(document, f"·  {item}", size=10.5, space_after=2)
        if meaning.watch:
            _paragraph(document, f"그래서 보라는 것  ·  {meaning.watch}", size=10.5, space_before=2, space_after=3)
        if meaning.grounds:
            _paragraph(document, "숫자로 받친 말  ·  검증하지 않음", size=10, bold=True, color=GRAY, space_before=2, space_after=2)
            for item in meaning.grounds:
                _paragraph(document, f"·  {item}", size=9.5, color=GRAY, space_after=1)
        meta = []
        if meaning.names:
            meta.append("이름  " + ", ".join(meaning.names))
        if meaning.sources:
            meta.append("시각  " + ", ".join(meaning.sources))
        if meta:
            _paragraph(document, "  ·  ".join(meta), size=8.5, color=GRAY, space_before=2, space_after=2)
    if briefing.agenda:
        _paragraph(document, "화자가 짚은 일정", size=14, bold=True, color=NAVY, space_before=12, space_after=4)
        for item in briefing.agenda:
            _paragraph(document, f"·  {item.what}", size=10.5, space_after=2)
    if briefing.basket:
        _paragraph(document, "화자가 적어 둔 관심 바구니", size=14, bold=True, color=NAVY, space_before=12, space_after=4)
        table = document.add_table(rows=1, cols=2)
        _shade(table.rows[0].cells[0], NAVY_HEX)
        _shade(table.rows[0].cells[1], NAVY_HEX)
        _cell(table.rows[0].cells[0], "묶음", bold=True, color=WHITE, size=9)
        _cell(table.rows[0].cells[1], "이름", bold=True, color=WHITE, size=9)
        for name, names in briefing.basket:
            row = table.add_row()
            _shade(row.cells[0], LIGHT_HEX)
            _cell(row.cells[0], name, bold=True, size=9)
            _cell(row.cells[1], names, size=9)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def write_outputs(briefing: Briefing, out_dir: Path, title: str) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    markdown_path = out_dir / "화자의 뜻.md"
    json_path = out_dir / "화자의 뜻.json"
    docx_path = out_dir / "화자의 뜻.docx"
    markdown_path.write_text(render_markdown(briefing, title), encoding="utf-8")
    json_path.write_text(render_json(briefing, title), encoding="utf-8")
    render_docx(briefing, title, docx_path)
    return {"md": markdown_path, "json": json_path, "docx": docx_path}


def _meaning_md(meaning: Meaning) -> list[str]:
    lines = [
        f"## {meaning.label}",
        "",
        f"**말하려는 것.** {meaning.point}",
        "",
    ]
    if meaning.against:
        lines.append(f"밀어낸 해석: {meaning.against}")
        lines.append("")
    if meaning.alongside:
        lines.append("같이 실린 판단")
        lines.append("")
        for item in meaning.alongside:
            lines.append(f"- {item}")
        lines.append("")
    if meaning.watch:
        lines.append(f"**그래서 보라는 것.** {meaning.watch}")
        lines.append("")
    if meaning.grounds:
        lines.append("숫자로 받친 말. 이 숫자를 맞추거나 고치지 않았다.")
        lines.append("")
        for item in meaning.grounds:
            lines.append(f"- {item}")
        lines.append("")
    meta = []
    if meaning.names:
        meta.append("이름 " + ", ".join(meaning.names))
    if meaning.sources:
        meta.append("시각 " + ", ".join(meaning.sources))
    if meta:
        lines.append(" · ".join(meta))
        lines.append("")
    return lines


def _first(text: str) -> str:
    import re

    parts = [part.strip() for part in re.split(r"(?<=다\.)\s+|(?<=요\.)\s+", text) if part.strip()]
    if len(parts) > 1 and len(parts[0]) < 42:
        return f"{parts[0]} {parts[1]}"
    return parts[0] if parts else text


def _setup(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(16)
    normal = document.styles["Normal"]
    normal.font.name = KR_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("퀵 코멘트  ·  화자의 뜻")
    _font(run, 8.5, False, GRAY)
    footer = section.footer
    footer.is_linked_to_previous = False
    foot = footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run("주장의 재구성  ·  사실 확인이 아님")
    _font(run, 8, False, GRAY)
    core = document.core_properties
    core.title = title
    core.author = "준혁"
    core.subject = "퀵 코멘트에서 화자의 판단을 증류"


def _paragraph(document, text, size=11, bold=False, color=DARK, space_after=6, space_before=0):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    _font(run, size, bold, color)
    return paragraph


def _callout(document, label: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    _shade(cell, AMBER_HEX)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(label + "  ")
    _font(run, 10, True, GOLD)
    run = paragraph.add_run(text)
    _font(run, 11, False, DARK)
    _set_cell_margins(cell)


def _cell(cell, text, bold=False, color=DARK, size=10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    _font(run, size, bold, color)
    _set_cell_margins(cell)


def _shade(cell, fill: str) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _set_cell_margins(cell) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="60" w:type="dxa"/>'
            f'<w:left w:w="90" w:type="dxa"/>'
            f'<w:bottom w:w="60" w:type="dxa"/>'
            f'<w:right w:w="90" w:type="dxa"/>'
            f"</w:tcMar>"
        )
    )


def _font(run, size, bold, color) -> None:
    run.font.name = KR_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
