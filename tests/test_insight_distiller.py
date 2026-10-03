"""추출·증류가 원문의 핵심 숫자와 테마를 보존하는지 확인한다."""

import re
from pathlib import Path

from docx import Document

from insight_distiller.distill import build_report
from insight_distiller.parse import dedupe_quick, near_duplicate, parse
from insight_distiller.render_docx import render_docx
from insight_distiller.render_md import render_markdown

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "paste_copy_2.txt"


def _report():
    return build_report(parse(SOURCE.read_text(encoding="utf-8")))


def test_parse_and_dedupe_quick_comments():
    blocks = parse(SOURCE.read_text(encoding="utf-8"))
    quick = [b for b in blocks if b.kind == "quick"]
    kept, dropped = dedupe_quick(quick)
    assert len(quick) >= 45
    assert dropped >= 4
    assert sum(b.kind == "chapter" for b in blocks) >= 4
    assert len(kept) < len(quick)
    assert any(b.kind == "chapter" for b in blocks)
    assert any(b.kind == "transcript" for b in blocks)


def test_near_duplicate_revision_collapses():
    a = "삼성전기투자 수혜 이익전망 기존 시장컨센대비 +64% 안전마진 20~30%까지는 오케"
    b = "삼성전기투자 수혜 이익전망 기존 시장컨센대비 +64% 안전마진 50~60%까지는 오케"
    assert near_duplicate(a, b)


def test_memory_hdd_and_portfolio_facts():
    scalars = _report().facts.scalars
    assert scalars["hynix_price"] == "184.2"
    assert scalars["hynix_per_26"] == "5.3"
    assert scalars["hynix_per_27"] == "4.1"
    assert scalars["hynix_op_26"] == "265"
    assert scalars["hynix_op_27"] == "392"
    assert scalars["samsung_price"] == "27.6"
    assert scalars["samsung_per_26"] == "5.8"
    assert scalars["samsung_per_27"] == "4.1"
    assert scalars["samsung_op_26"] == "388"
    assert scalars["micron_fy_eps"] == "165"
    assert scalars["micron_fy_per"] == "6.5"
    assert scalars["sandisk_per"] == "7.5"
    assert scalars["stx_per"] == "26"
    assert scalars["wdc_per"] == "23"
    assert scalars["jobs_actual"] == "2.9"
    assert scalars["jobs_cons"] == "8.9"
    assert scalars["port_ai"] == "50~60%"
    assert scalars["port_cash"] == "20%"
    assert scalars["port_cash_light"] == "30%"
    assert scalars["semco_margin"] == "20~30%"
    assert scalars["sanil_per"] == "16"
    assert scalars["intech_tp"] == "9.5"


def test_report_sections_and_headline():
    report = _report()
    keys = {section.key for section in report.sections}
    assert {"macro", "memory", "hdd", "optical", "substrate", "power", "portfolio"} <= keys
    assert "싸움" in report.headline or "장기금리" in report.headline
    assert [section.key for section in report.sections][:2] == ["macro", "memory"]
    assert report.facts.scalars["floor_per"] == "7"
    assert report.facts.scalars["hynix_floor"] == "209~244"
    portfolio = next(section for section in report.sections if section.key == "portfolio")
    assert "30%%" not in portfolio.lead
    assert all("모시고" not in bullet.text for bullet in portfolio.bullets)
    risk = next(section for section in report.sections if section.key == "risk")
    assert all("레버리지 써서" not in bullet.text for bullet in risk.bullets)
    assert all(
        not re.search(r"\d{1,2}:\d{2}", bullet.text)
        for section in report.sections
        for bullet in section.bullets
    )
    memory = next(section for section in report.sections if section.key == "memory")
    assert memory.table and memory.table[0]["name"] == "SK하이닉스"
    assert memory.lead
    assert memory.bullets
    markdown = render_markdown(report)
    assert "SK하이닉스" in markdown
    assert "5.3" in markdown
    assert "매수·매도 권유가 아니다" in markdown


def test_docx_contains_distilled_figures(tmp_path):
    report = _report()
    path = tmp_path / "insights.docx"
    render_docx(report, path)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " ".join(cell.text for cell in row.cells)
    assert "투자 인사이트" in text
    assert "184.2만 원" in text
    assert "PER 5.3배" in text
    assert doc.core_properties.author == "준혁"
