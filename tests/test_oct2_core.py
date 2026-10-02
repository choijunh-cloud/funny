"""핵심 메모가 포지션 규칙을 빠뜨리지 않는지."""

from __future__ import annotations

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document  # noqa: E402

from oct2_core import OUT_PATH, build  # noqa: E402


class Oct2CoreTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        build()
        doc = Document(str(OUT_PATH))
        cls.text = "\n".join(p.text for p in doc.paragraphs)
        cls.text += "\n" + "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )

    def test_core_keeps_the_book_and_the_breaks(self):
        for snippet in [
            "24만 5,000원",
            "20~30%",
            "3.008%",
            "1,500억",
            "6,971",
            "420억",
            "1350",
            "소비와 GDP",
            "매수·매도 권유 아님",
        ]:
            self.assertIn(snippet, self.text, snippet)


if __name__ == "__main__":
    unittest.main()
