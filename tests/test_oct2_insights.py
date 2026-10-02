"""원문 앵커와 브리프가 어긋나지 않는지 확인한다."""

from __future__ import annotations

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document  # noqa: E402

from oct2_analyze import (  # noqa: E402
    FACT_NEEDLES,
    check_facts,
    find_tensions,
    load_corpus,
    portfolio_books,
)
from oct2_brief import OUT_PATH, build  # noqa: E402


class Oct2InsightsTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.corpus = load_corpus()
        cls.payload = build()
        cls.doc = Document(str(OUT_PATH))
        cls.text = "\n".join(p.text for p in cls.doc.paragraphs)
        cls.text += "\n" + "\n".join(cell.text for table in cls.doc.tables for row in table.rows for cell in row.cells)

    def test_every_fact_needle_is_in_the_source(self):
        missing = [row["id"] for row in check_facts(self.corpus) if not row["found"]]
        self.assertEqual(missing, [])
        self.assertGreaterEqual(len(FACT_NEEDLES), 40)

    def test_portfolio_ranges_match_the_comment(self):
        books = portfolio_books()
        low, high = books["experienced_ai50"], books["experienced_ai60"]
        self.assertEqual(sum(low.values()), 100)
        self.assertEqual(sum(high.values()), 100)
        self.assertEqual(low["삼전닉스"] + low["소부장"], 50)
        self.assertEqual(high["삼전닉스"] + high["소부장"], 60)
        self.assertEqual(low["현금"], 20)
        self.assertEqual(low["2차전지"], 10)
        self.assertEqual(low["원문 미배분"], 10)
        self.assertEqual(high["원문 미배분"], 0)

    def test_tensions_include_the_two_sided_calls(self):
        ids = {row["id"] for row in find_tensions(self.corpus)}
        self.assertIn("semco_cushion", ids)
        self.assertIn("hike_cycle", ids)
        self.assertIn("muse_two_sided", ids)
        self.assertIn("krw_2027", ids)

    def test_brief_keeps_the_operating_rules(self):
        for snippet in [
            "왜 고점인지",
            "3.344",
            "2.983",
            "3.008",
            "86.2%",
            "1,500억",
            "24만 5,000원",
            "20~30%",
            "50~60%",
            "420억",
            "에어비앤비",
            "산일전기",
            "인텍플러스",
            "매수·매도 권유가 아니다",
            "원문 미배분",
        ]:
            self.assertIn(snippet, self.text, snippet)

    def test_brief_embeds_the_charts(self):
        self.assertGreaterEqual(len(self.doc.inline_shapes), 8)
        self.assertGreaterEqual(len(self.doc.tables), 20)
        charts = Path(ROOT / "output" / "oct2" / "charts")
        for name in [
            "theme_density.png",
            "ust_move.png",
            "scoreboard.png",
            "portfolio.png",
            "pce_path.png",
            "scenario_map.png",
            "chapters.png",
            "capex_timeline.png",
        ]:
            path = charts / name
            self.assertTrue(path.exists(), name)
            self.assertGreater(path.stat().st_size, 8_000, name)


if __name__ == "__main__":
    unittest.main()
