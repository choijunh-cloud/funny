"""화자의 판단이 숫자 표보다 앞에 오는지 고정한다."""

from pathlib import Path

from docx import Document

from insights.distill import distill_path, distill_text
from insights.intent import factiness
from insights.parse import parse_comments
from insights.render import render_docx, render_markdown

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "data" / "raw" / "2026-09-29_quick_comments.txt"


def test_glued_timestamps_stay_with_the_previous_comment():
    comments = parse_comments(SAMPLE.read_text(encoding="utf-8"))
    by_time = {}
    for comment in comments:
        by_time.setdefault(comment.time, comment)
    assert comments[0].time == "20:34"
    assert len(comments) >= 50
    assert "SKHY -5.03%" in by_time["07:14"].body
    assert "SKHY +2.62%" not in by_time["07:14"].body
    assert "SKHY +2.62%" in by_time["07:47"].body
    assert "오늘 지수 반등" in by_time["15:36"].body


def test_short_verdict_outranks_a_per_table():
    text = """
08:15
Quick 코멘트
괴리 너무 큰 것으로 봅니다
08:15
Quick 코멘트
SK하이닉스 본주 PER 4.0배
삼성전자 PER 4.05배
마이크론 PER 7.1배
할인 괴리가 핵심입니다.
"""
    briefing = distill_text(text)
    meaning = briefing.by("memory_valuation")
    assert "괴리" in meaning.point
    assert not meaning.point.startswith("SK하이닉스")
    assert factiness(meaning.point) < 0.12


def test_sample_day_follows_what_the_speaker_is_arguing():
    briefing = distill_path(SAMPLE)
    macro = briefing.by("macro_rates")
    memory = briefing.by("memory_valuation")
    samsung = briefing.by("samsung_dispersion")
    micron = briefing.by("micron_print")
    hbm = briefing.by("hbm")
    supply = briefing.by("ai_supply")
    demand = briefing.by("ai_demand")
    tesla = briefing.by("tesla_fsd")
    method = briefing.by("method")
    sdi = briefing.by("sdi")

    assert "결과" in macro.point
    assert "원인" in macro.against
    assert "마이크론" not in macro.point

    assert "괴리" in memory.point
    assert any("PER" in ground or "배" in ground for ground in memory.grounds)

    assert "2028" in samsung.point or "초호황" in samsung.point
    assert "346" not in samsung.point
    assert any("4Q" in item or "가격" in item for item in [samsung.watch, *samsung.alongside])

    assert "마이크론" in micron.point
    assert any(word in micron.point + micron.watch + " ".join(micron.alongside) for word in ("방향", "핵심", "신뢰"))
    assert any("가이던스" in ground or "GPM" in ground for ground in micron.grounds)

    assert "후공정" in hbm.point or "증가" in hbm.point
    assert "동의" in hbm.point or "동감" in hbm.point
    assert "수요 감소" in hbm.against

    assert "ABF" in supply.point or "기판" in supply.point
    assert "싸게" in demand.point or "오래" in demand.point
    assert "기대" in tesla.point
    assert "실패" in tesla.against
    assert "커버" in method.point or "수주" in method.point
    assert "흑자" in sdi.point or "유리" in sdi.point

    blob = " ".join(meaning.point for meaning in briefing.meanings)
    assert "SKIP" not in blob
    for meaning in briefing.meanings:
        assert factiness(meaning.point) < 0.15


def test_render_leads_with_meaning_and_labels_numbers_as_unverified(tmp_path):
    briefing = distill_path(SAMPLE)
    title = "9월 29일 퀵 코멘트, 화자가 말한 뜻"
    markdown = render_markdown(briefing, title)
    assert markdown.index("말하려는 것") < markdown.index("숫자로 받친 말")
    assert "맞추거나 반박하지 않았다" in markdown
    assert "검증하지 않음" in markdown or "맞추거나 고치지 않았다" in markdown

    docx_path = tmp_path / "brief.docx"
    render_docx(briefing, title, docx_path)
    document = Document(docx_path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    for section in document.sections:
        chunks.extend(paragraph.text for paragraph in section.footer.paragraphs)
    text = "\n".join(chunks)
    assert "말하려는 것" in text
    assert "사실 확인이 아님" in text
    assert "괴리" in text
