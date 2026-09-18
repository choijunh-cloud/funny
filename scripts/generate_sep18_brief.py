#!/usr/bin/env python3
"""9월 18일 AI·반도체 시장 코멘트 강의노트(.docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor

OUT_PATH = Path("/workspace/lectures/9월 18일 AI·반도체 시장 코멘트.docx")
CHARTS = Path("/workspace/lectures/assets/sep18")

KR_FONT = "맑은 고딕"
NAVY = RGBColor(0x0F, 0x20, 0x43)
NAVY2 = RGBColor(0x1E, 0x40, 0x7C)
GOLD = RGBColor(0xB8, 0x94, 0x3A)
GRAY = RGBColor(0x4B, 0x55, 0x63)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x16, 0x65, 0x34)
RED = RGBColor(0x99, 0x1B, 0x1B)
AMBER = RGBColor(0x7A, 0x5C, 0x12)

NAVY_HEX, NAVY2_HEX, GOLD_HEX = "0F2043", "1E407C", "B8943A"
LIGHT_HEX, GREEN_HEX, RED_HEX = "EEF2F8", "E8F5E9", "FDECEA"
AMBER_HEX, BLUE_HEX, ROW_HEX, WHITE_HEX = "FFF8E7", "E8F1FB", "F7F9FC", "FFFFFF"


def set_run_font(run, size=11, bold=False, color=DARK, italic=False, font=KR_FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/>'
            f"</w:tcMar>"
        )
    )


def set_table_borders(table, color="D0D7E2", sz="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    tbl_pr.append(
        parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f"</w:tblBorders>"
        )
    )


def set_left_accent(cell, color=NAVY_HEX, sz="24"):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="nil"/><w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="nil"/><w:right w:val="nil"/>'
            f"</w:tcBorders>"
        )
    )


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def cell_text(cell, text, size=10, bold=False, color=DARK, align="left"):
    cell.text = ""
    align_enum = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = align_enum
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1 if i < len(lines) - 1 else 0)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(line)
        set_run_font(run, size=size, bold=bold, color=color)
    set_cell_margins(cell)


def add_runs(paragraph, parts, size=11, color=DARK):
    for part in parts:
        if isinstance(part, str):
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)
        else:
            text, bold, *rest = part
            c = rest[0] if rest else color
            run = paragraph.add_run(text)
            set_run_font(run, size=size, bold=bold, color=c)


class Notes:
    def __init__(self):
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width = Mm(210)
        sec.page_height = Mm(297)
        sec.left_margin = Mm(16)
        sec.right_margin = Mm(16)
        sec.top_margin = Mm(16)
        sec.bottom_margin = Mm(16)
        sec.header_distance = Mm(8)
        sec.footer_distance = Mm(8)

        normal = self.doc.styles["Normal"]
        normal.font.name = KR_FONT
        normal.font.size = Pt(11)
        normal.font.color.rgb = DARK
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.18

        header = sec.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("9/18  AI·반도체 시장 코멘트  ·  BOJ 인상일  ·  강의노트")
        set_run_font(r, size=8.5, color=GRAY)

        footer = sec.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("영상 5편 + 퀵코멘트를 주제별로 재구성  ·  숫자는 공개 자료로 교차검증  ·  ")
        set_run_font(r, size=8, color=GRAY)
        fld = parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE ">'
            f'<w:r><w:rPr><w:sz w:val="16"/><w:color w:val="4B5563"/>'
            f'<w:rFonts w:ascii="{KR_FONT}" w:hAnsi="{KR_FONT}" w:eastAsia="{KR_FONT}"/></w:rPr>'
            f"<w:t></w:t></w:r></w:fldSimple>"
        )
        fp._p.append(fld)

        core = self.doc.core_properties
        core.title = "9월 18일 AI·반도체 시장 코멘트"
        core.author = "준혁"
        core.subject = "BOJ, 노드, 메모리, Solidigm, Generac, 속도조절"

    def p(self, text, size=11, bold=False, color=DARK, space_after=6, space_before=0, align="left"):
        para = self.doc.add_paragraph()
        para.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(space_before)
        para.paragraph_format.line_spacing = 1.18
        run = para.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
        return para

    def h1(self, text, num=None):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(14)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.line_spacing = 1.1
        if num:
            run = para.add_run(f"{num}  ")
            set_run_font(run, size=16, bold=True, color=GOLD)
        run = para.add_run(text)
        set_run_font(run, size=16, bold=True, color=NAVY)
        para._p.get_or_add_pPr().append(
            parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="12" w:space="4" w:color="{NAVY_HEX}"/>'
                f"</w:pBdr>"
            )
        )
        return para

    def h2(self, text):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(4)
        run = para.add_run(text)
        set_run_font(run, size=13, bold=True, color=NAVY2)
        return para

    def bullet(self, text, level=0, bold_lead=None, size=11):
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.55 + level * 0.45)
        para.paragraph_format.first_line_indent = Cm(-0.35)
        para.paragraph_format.space_after = Pt(2.5)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing = 1.15
        mark = "• " if level == 0 else "– "
        run = para.add_run(mark)
        set_run_font(run, size=size, color=NAVY2 if level == 0 else GRAY)
        if bold_lead:
            run = para.add_run(bold_lead)
            set_run_font(run, size=size, bold=True, color=DARK)
            run = para.add_run(text)
            set_run_font(run, size=size, color=DARK)
        else:
            run = para.add_run(text)
            set_run_font(run, size=size, color=DARK)
        return para

    def flow(self, items, size=11):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.line_spacing = 1.2
        for i, item in enumerate(items):
            if i:
                run = para.add_run("   →   ")
                set_run_font(run, size=size, bold=True, color=GOLD)
            run = para.add_run(item)
            set_run_font(run, size=size, bold=True, color=NAVY)
        return para

    def callout(self, title, body, kind="key"):
        palette = {
            "key": (NAVY_HEX, LIGHT_HEX, NAVY),
            "bull": ("166534", GREEN_HEX, GREEN),
            "bear": ("991B1B", RED_HEX, RED),
            "note": (GOLD_HEX, AMBER_HEX, AMBER),
            "blue": (NAVY2_HEX, BLUE_HEX, NAVY2),
        }
        accent, fill, title_color = palette[kind]
        table = self.doc.add_table(rows=1, cols=1)
        table.autofit = True
        table.allow_autofit = True
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        shade_cell(cell, fill)
        set_left_accent(cell, accent, sz="28")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.space_before = Pt(0)
        r = p1.add_run(title)
        set_run_font(r, size=10, bold=True, color=title_color)
        if isinstance(body, str):
            body = [body]
        for line in body:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(line)
            set_run_font(r, size=10.5, color=DARK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def table(self, headers, rows, col_widths=None, first_col_bold=True):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table, color="D5DCE6", sz="4")
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            shade_cell(cell, NAVY_HEX)
            cell_text(cell, h, size=9.5, bold=True, color=WHITE, align="center")
        prevent_row_split(table.rows[0])
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = table.rows[r_i + 1].cells[c_i]
                shade_cell(cell, ROW_HEX if r_i % 2 else WHITE_HEX)
                align = "left" if c_i == 0 else "center"
                bold = first_col_bold and c_i == 0
                cell_text(cell, str(val), size=9.5, bold=bold, color=DARK, align=align)
            prevent_row_split(table.rows[r_i + 1])
        if col_widths:
            for row in table.rows:
                for i, w in enumerate(col_widths):
                    row.cells[i].width = Cm(w)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)
        spacer.paragraph_format.space_before = Pt(2)
        return table

    def image(self, name, width=16.4):
        path = CHARTS / name
        if not path.exists():
            self.p(f"[차트 없음: {name}]", size=9, color=GRAY)
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width))

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))


def build():
    n = Notes()

    n.p("2026. 9. 18. 강의노트  ·  FOMC 다음날 + BOJ 인상일  ·  영상 5편 + 퀵코멘트", size=10.5, color=GRAY, align="center", space_after=4)
    n.p("금리 산은 넘고 노드는 남았다", size=13, bold=True, color=GOLD, align="center", space_after=2)
    n.p("9월 18일  AI·반도체 시장 코멘트", size=22, bold=True, color=NAVY, align="center", space_after=2)
    n.p("노드  ·  메모리 2031  ·  미국 3축  ·  Generac  ·  속도조절", size=12, bold=True, color=NAVY2, align="center", space_after=8)

    n.callout(
        "한 장 결론",
        [
            "연준·BOJ가 올린 뒤에도 주식은 안도했다. 알려진 인상은 악재가 아니었다.",
            "정지훈의 프레임: 병목=노드=돈. 연산은 풀렸고 지금은 메모리. 에너지는 가장 복제하기 어렵다.",
            "Amazon–Generac $24억(최대 $80억)이 전력 캡엑스의 영수증이다. 전선주는 장중 급등, 밸류는 이미 비싸다.",
            "하이닉스 미국은 Indiana(후공정 확정) · Ohio(협의) · New York Solidigm(검토) 세 줄. 확정이 아니다.",
            "Citi는 지속학습이 메모리 부족을 2031까지 연장할 수 있다고 본다. 젠슨의 칩 2배는 대수지, 매출 2배가 아니다.",
        ],
        kind="key",
    )

    n.h2("이 노트를 읽는 법")
    n.table(
        ["표시", "뜻"],
        [
            ["사실", "거래소, BOJ 성명, Reuters, Citi, 8-K, CNBC로 확인된 것"],
            ["부분 / 주의", "방향은 맞지만 숫자·시제·범위가 다른 것"],
            ["해석", "방송 논리로 쓸 수는 있으나 확정처럼 쓰면 안 되는 것"],
            ["쓰지 말 것", "Ohio/NY 투자 확정, 젠슨 2배=매출 2배, 원/달러 1,587, 매일 1.6조 자사주, CXMT 즉각 공급 해소"],
        ],
        col_widths=[3.6, 14.0],
    )

    n.h2("오늘 재료가 온 곳")
    n.table(
        ["소스", "핵심"],
        [
            ["정지훈 박사 대담", "문명 가속기, 플라이휠, 노드=병목=돈, 디지털+피지컬"],
            ["증시각도기 9/18", "KOSPI 6,894, BOJ 7-2, 엔 약세, Generac, 속도조절 기구 무산"],
            ["문남중 (대신)", "불확실성 제거, 실질정책금리, 엔캐리 과장, 추석 전후"],
            ["박승진 (하나)", "점도표 < 시장, 12월 인상, IT ETF, 부외리스, 내년 이익 둔화"],
            ["빈센트 김두원", "패러다임, 5% 집착 말 것, 시대 주식=AI, CXMT 옛 문법 아님"],
            ["증시입담화", "산 하나 넘음, 투톱 홀드, 소부장 ETF, 21일 전력 특집"],
            ["퀵코멘트", "Solidigm NY, CXMT NAND, Citi 2031, 젠슨 2배, DOE/PJM"],
        ],
        col_widths=[4.4, 13.2],
    )

    n.h1("오늘 숫자 한 장", num="0.")
    n.table(
        ["항목", "숫자", "한 줄"],
        [
            ["KOSPI / 코스닥", "6,894.23 (+2.66%)\n827.12 (+0.60%)", "6,900은 장중만. 시총 상단이 밀었다"],
            ["삼성 / 하이닉스", "261,000 (+3.37%)\n1,857,000 (+6.42%)", "외인은 닉스 +1.33조, 삼성 −0.46조"],
            ["수급 (KRX)", "개인 −3.59조\n기관 +1.50 / 법인 +1.67", "추석 전 개인 차익실현"],
            ["BOJ", "1.25%  ·  7–2", "31년 만의 고점. 엔은 157로 약세"],
            ["환율 / 금리", "원/달러 1,383.3\n국고 10년 4.466%(−4bp)", "방송 1,587은 오역"],
            ["젠슨 9/17", "칩 대수 내년 2배\nFY28 매출 +70%", "대수 ≠ 매출"],
            ["Generac", "$24억 (27–28)\n한도 $80억", "비상발전기. UPS 아님"],
            ["Citi 메모리", "DRAM 부족 8.7→9.7%\nNAND 6.1→5.5%", "지속학습 → 2031까지 연장 가능"],
        ],
        col_widths=[3.4, 6.2, 8.0],
    )

    n.h1("시장 — 산은 넘고 개인은 팔았다", num="1.")
    n.image("01_market_rebound.png")
    n.image("02_flows.png")
    n.bullet("KOSPI 6,894.23, +178.82(+2.66%). 고가 6,914.08. 오른 종목보다 내린 종목이 많았다. 대형 전기전자가 지수를 끌었다.")
    n.bullet("9/17 미국: 다우 +0.6, S&P +1.1, 나스닥 +1.7, SOX +3.1. 전날 10년 5% 부담의 기술적 되돌림 + 유가 하락.")
    n.bullet("개인 −3.59조, 기관 +1.50조, 외인 +0.43조, 기타법인 +1.67조(자사주). NXT 합산 시 개인 −4.37조.")
    n.bullet("방송 “매일 평균 1.6조 자사주”는 당일 기타법인과 혼동된 표현으로 본다. 당일 1.67조는 사실, 매일 평균은 아니다.")
    n.callout("추석", ["문남중: 최근 10년 추석 전 5일 KOSPI −0.42%, 후 +0.68%. 입담화: 급하지 않으면 연휴 전에 팔지 말라."], kind="note")

    n.h1("BOJ — 올렸는데 엔은 약하다", num="2.")
    n.image("03_boj_yen.png")
    n.bullet("7–2로 1.00→1.25%. 9/24 시행. 반대는 아사다(신선식품 제외 CPI가 2%를 밑돈다), 사토(가속이 없다).")
    n.bullet("성명은 계속 올리겠다고 했지만 7월 문구와 큰 차이가 없다. 엔 157.12(+1.15). 원/엔 880.27.")
    n.bullet("원/달러 종가 1,383.3(+1.1원), 7거래일 연속 상승. 장중 고가 1,387.1. 방송의 1,587은 1,387 오역.")
    n.bullet("국고 10년 4.466%(−4.0bp). 방송 −2.5bp는 장중 스케치.")
    n.bullet("문남중: 미·일 10년 스프레드 약 2%p면 엔캐리 청산은 어렵다. 일본 기업의 미국 직접투자가 늘었다.")
    n.bullet("증시각도기: 31년 만의 고점이지만 글로벌 금리와 비교하면 여전히 낮다. 연타로 올리면 그때는 충격.")

    n.h1("노드와 플라이휠 — 정지훈", num="3.")
    n.image("04_nodes.png")
    n.image("09_flywheel.png")
    n.p("정지훈은 AI를 특정 IT 섹터가 아니라 문명 가속기(augmented intelligence)로 본다. 책은 종목 추천이 아니라 2025–2035 10년 지도다. 테마는 빨리 당기므로, 사인이 왔을 때 놓치지 않게 하는 것이 목표다.")
    n.h2("플라이휠")
    n.bullet("아마존 냅킨: 이익을 가격 인하·구색에 재투자 → 고객 경험 → 사용자 → 볼륨 → 다시 가격. 이익을 못 내는 게 아니라 안 낸다.")
    n.bullet("임계를 넘으면 마지막 노드인 캐피탈이 스스로 돈다. 성공하면 사업가, 실패하면 사기꾼으로 불린다.")
    n.bullet("PC–윈텔은 한 사이클. AI는 디지털(토큰)과 피지컬(로봇·공장)이 시차를 두고 동시에 돈다.")
    n.h2("토큰 장사")
    n.bullet("측정치(해석): 같은 토큰을 Anthropic이 OpenAI보다 약 5배 비싸게 받아도 사람들이 남는다. 명품 전략.")
    n.bullet("구글은 TPU·Flash·워크스페이스 번들로 원가를 낮춘다. 가격이 떨어지면 일을 더 시킨다. 수요가 죽는 그림이 아니다.")
    n.h2("노드")
    n.flow(["연산 풀림", "메모리 전 영역", "에너지", "지식/서비스", "자본"])
    n.bullet("연산 병목은 출구·대체재·하이퍼스케일러 자체 칩으로 많이 풀렸다. 엔비디아 이익률 정점이 꺾이는 신호가 그것이다.")
    n.bullet("지금은 어떤 연산 반도체를 만들어도 메모리가 필요하다. 에이전틱은 중간산출과 컨텍스트를 쌓는다. HBM이 아니어도 된다.")
    n.bullet("시뮬레이터에서 가장 심각한 병목은 에너지. 땅·물·전력 중 전력이 가장 복제하기 어렵다. 미국 주민 반대의 실체는 여름 정전 공포.")
    n.bullet("해법은 세 갈래. ①효율(메타 Bailey 광스위치 +40% → 루멘텀 → 코닝) ②생산(SMR이 기존 화력 인프라에 들어감) ③유연 부하(AEMA).")
    n.callout("핵심 문장", ["에너지가 풀리면 실리콘 병목이 다시 올라온다. 다음 견인차는 삼성·하이닉스만이 아닐 수 있다."], kind="key")

    n.h1("메모리 — 지속학습이 부족을 늘린다", num="4.")
    n.image("05_memory_gap.png")
    n.bullet("Citi: 지속학습은 새 과제를 배우면서 과거 데이터에도 접근한다. HBM·서버 DDR5·eSSD가 같이 는다. 부족은 2031까지 갈 수 있다.")
    n.bullet("HBM 비트 수요 +62%(27) / +69%(28). DRAM 수요 +30/+35 vs 공급 +19/+22 → 부족 8.7→9.7%. NAND +29/+33 vs +21/+25 → 6.1→5.5%.")
    n.bullet("선호 이름: 삼성, 하이닉스, 마이크론, 샌디스크, 키옥시아.")
    n.bullet("젠슨 9/17 스코틀랜드: 내년 칩 판매 대수가 올해의 두 배. 매출 가이던스 FY28 약 $673B(+70%)와는 별개. 제품 믹스가 있어 대수가 매출을 결정하지 않는다.")
    n.bullet("대만 매체 4Q DRAM +20% / NAND +30% QoQ. 퀵코멘트도 원소스가 불명확하다고 적어 두었다. 방향으로만.")
    n.bullet("문남중: 글로벌 반도체 매출 증가, 메모리 가격 4월 이후 상승, 한국 8월 반도체 수출 +209%·18개월 연속.")

    n.h1("미국 공급망 — 세 줄 + CXMT", num="5.")
    n.image("06_us_footprint.png")
    n.table(
        ["축", "날짜", "상태", "쓰지 말 것"],
        [
            ["Indiana HBM 후공정", "기존", "진행. $40억+ · 2029H2", "오늘 신규 뉴스처럼"],
            ["Ohio × Intel 메모리", "9/16 Reuters", "탐색. DRAM/NAND 미정", "계약·양산 확정"],
            ["NY Solidigm NAND", "9/18 Reuters", "부지 검토. 다롄 분산", "미국 Fab 투자 확정"],
            ["CXMT 베이징 NAND", "9/18 Reuters", "R&D·시험생산. 양산 시점 없음", "당장 공급 해소"],
        ],
        col_widths=[4.2, 3.2, 5.6, 4.6],
    )
    n.bullet("Solidigm 검토의 성격: ①다롄 의존 축소 ②엔터프라이즈 SSD 미국 현지화 ③CXMT NAND 진입 대응 ④미국 생산 압박 대응.")
    n.bullet("미국 NAND는 한국보다 원가가 높다. Reuters도 경제성·보조금·고객 LTA를 변수로 적었다.")
    n.bullet("빈센트: 과거 문법으로는 중국 진입=시장 붕괴. 지금은 비싼 물건이 팔린다. CXMT는 내수용 구획이 될 수 있다. 하방 압력은 보되 끝장으로 쓰지 말 것.")

    n.h1("전력 — 영수증이 나왔다", num="6.")
    n.image("08_generac.png")
    n.bullet("Amazon–Generac 8-K(9/16). 비상발전기 27–28년 초기 $24억, 워런트 베스팅 한도 $80억. 아마존 워런트 최대 169만주 @ $200.93.")
    n.bullet("방송이 UPS/USP로 헤맨 부분은 오해. 계약의 본체는 산업용 백업 제너레이터다.")
    n.bullet("가온전선 장중 +18~22%, 고가 326,000. 미국 LSCUS 증설(생산능력 2배+)이 겹쳤다. 입담화: 가온은 이미 비싸다. LS·일진·산일·현대일렉·효성 쪽에서 덜 달린 것을 보라.")
    n.bullet("DOE/PJM: 9/16 최대발전·부하관리 경보. 백업발전 동원 권한 신청. 확인된 명령 202-26-41은 9/1–8. 9/17–18 한시 명령은 퀵코멘트·신청 보도 기준.")
    n.callout("연결", ["캡엑스가 GPU에서 전선·변압기·비상전원으로 번진다. AEMA는 유연 부하이지 전력기기 수요의 소멸이 아니다."], kind="blue")

    n.h1("한국 성장과 환율", num="7.")
    n.image("07_korea_gdp.png")
    n.bullet("Citi 9/18: 실질 GDP 3.7 / 3.1 / 3.0. 한은 8월 3.3 / 2.9보다 위. 반도체 수출 +175% / +41%.")
    n.bullet("올해 실질 기여 +3.2%p, 명목 +15.6%p. 명목 GDP 25%는 1981년 이후 최고 가정.")
    n.bullet("문남중: 원화가 1,300초까지 내려온 구간은 삼전·닉스 환차손 3–5조 부담. 전쟁·달러 강세가 1,380–1,400으로 되돌리고 있다.")
    n.bullet("빈센트: 펀더멘털은 아래, 심리는 위 → 횡보. 1,400을 다시 넘어도 예전처럼 1,500을 보는 공포는 덜할 것. 환율 영향은 3분기 실적 발표 후 많이 해소.")
    n.bullet("입담화: 대미 투자 달러 수요 + 연준 인상 = 원화 약세가 수출 실적에는 오히려 숨통.")

    n.h1("속도조절 · 중국 · 정치", num="8.")
    n.bullet("WSJ: 젠슨·저커버그·머스크가 업계 자율규제기구를 막았다. 허사비스식 FINRA 모델. 백악관 최소규제. 주정부·전력 인허가는 남는다.")
    n.bullet("문남중·빈센트: 속도조절은 노이즈. 말한 회사들은 속도를 안 줄인다. 중국은 같이 멈춰 주지 않는다.")
    n.bullet("빈센트: 민간이 정부에 규제를 요청하는 역전. 중간선거의 아젠다가 주식에서 데이터센터로 이동. 선거 후 민주당도 중국을 이유로 말을 바꿀 수 있다.")
    n.bullet("미중 정상회담: 빈센트는 당장 소득 없음, 다만 충돌 전 골든타임. 미국은 장기금리를 누르고 싶고 중국은 투자처가 없다(10년 5% vs 2%). 중국의 선물은 ‘매도 일시 정지’ 정도.")
    n.bullet("입담화: 둘이 너무 친해지면 한국 반사익(2차전지·통신장비·태양광, 장비 수출통제)이 줄어든다.")
    n.bullet("골드만 중국 가격 모델: 싸게 깔고 점유 후 올린다. 로봇청소기·변압기·H형강. 전기차·로봇택시는 아직 내부 출혈.")
    n.bullet("모건스탠리 중국 산업 전환 투자. 방송 15조 vs 본문 10조 달러 — 숫자 혼선. 방향으로만.")

    n.h1("전문가 맵과 포트", num="9.")
    n.table(
        ["사람", "사면 되는가", "무기로 보는 것", "찜찜한 것"],
        [
            ["문남중", "불확실성 제거 후 안도", "HBM·전력·DC. 한·미 반도체 같이", "추석 전 경계, 가상자산은 3년"],
            ["박승진", "지금 표면 지표로는 매수", "XLK·IYW·VGT·BAI·CHAT", "12월 인상, 부외리스, 내년 이익 둔화, 소비 바닥"],
            ["빈센트", "시대 주식을 모은다", "한국 1·2등 하드웨어", "CXMT 하방, 외인 시차 6개월"],
            ["입담화", "연휴 전 매도 아님", "투톱 홀드 30–40%, 소부장 10–20%", "가온 과열, 돈이 없어 업종 로테이션"],
            ["증시각도기", "느긋하게", "반도체+전력", "28년 긴축 시차, 평균회귀 경고"],
        ],
        col_widths=[2.8, 4.2, 5.4, 5.2],
    )
    n.bullet("박승진: 반도체는 최전방 전사라 너덜너덜하다. 코어로 두면 변동성이 너무 크다. 에너지 업종은 한 번 뒤집히면 2–3년 버린다.")
    n.bullet("박승진이 전한 부외리스: 재무제표 부채 약 $1.7조, 숨은 것까지 $3조. 운영리스. 월스트리트저널 이슈. 언제든 재점화.")
    n.bullet("경고 리포트는 컨센이 아니다. NZ Super는 20년 평균의 두 배 수익률 → 평균회귀. BofA는 CAPE 32면 10년 연 −3% 가능. 핵심 반론은 AI 성장이 그 회귀를 늦추느냐다.")
    n.bullet("퀵코멘트 원칙: 내러티브만 쫓는 종목은 higher for longer에서 위험. 소부장 아웃퍼폼의 전제는 삼전닉스가 폭발하지 않고 안정적일 것, 그리고 종목별 해자가 있을 것.")

    n.h1("고친 것 · 확인 넷", num="10.")
    n.table(
        ["방송/코멘트", "고친 값", "이유"],
        [
            ["원/달러 1,587", "1,383.3 (장중 1,387)", "1,387 오역"],
            ["국고 10년 −2.5bp", "종가 −4.0bp / 4.466%", "장중 스케치"],
            ["매일 1.6조 자사주", "당일 기타법인 +1.67조", "평균이 아님"],
            ["젠슨 칩 2배", "대수. 매출 +70%는 FY28", "믹스 때문에 불일치"],
            ["Generac = UPS", "비상발전기 8-K", "용어 혼선"],
            ["NY NAND 확정", "검토, 결정 없음", "Reuters·회사 코멘트"],
            ["CXMT 양산", "R&D·시험, 시점 없음", "같은 날 별도 뉴스"],
            ["Anthropic $2조 IPO", "내부 전언 수준", "사실로 쓰지 말 것"],
        ],
        col_widths=[4.6, 6.4, 6.6],
    )
    n.callout(
        "가을 확인 넷",
        [
            "10/1 마이크론 마진이 80%대 후반 스토리를 받쳐 주는지.",
            "9/29 DevDay에서 에이전트·API 토큰·기업 배포가 숫자로 나오는지.",
            "Ohio·New York이 탐색에서 계약으로 넘어가는지. Indiana는 이미 별축.",
            "10년 5%와 WTI $100–110이 다시 붙는지. 붙으면 같은 조정 논리가 반복된다.",
        ],
        kind="key",
    )
    n.p("다음 방송 힌트: 입담화는 9/21(월) 13시 ‘AI 데이터센터발 병목, 전력인가 전력망인가’ 특집을 예고했다.", size=10.5, color=GRAY)

    n.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({OUT_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
