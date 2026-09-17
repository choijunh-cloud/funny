#!/usr/bin/env python3
"""9월 17일 AI·반도체 시장 코멘트 강의노트(.docx) 생성."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

OUT_PATH = Path("/workspace/lectures/9월 17일 AI·반도체 시장 코멘트.docx")
CHARTS = Path("/workspace/lectures/assets/sep17")

KR_FONT = "맑은 고딕"
NAVY = RGBColor(0x0F, 0x20, 0x43)
NAVY2 = RGBColor(0x1E, 0x40, 0x7C)
GOLD = RGBColor(0xB8, 0x94, 0x3A)
GRAY = RGBColor(0x4B, 0x55, 0x63)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x16, 0x65, 0x34)
RED = RGBColor(0x99, 0x1B, 0x1B)
AMBER = RGBColor(0x7A, 0x5C, 0x12)

NAVY_HEX, NAVY2_HEX, GOLD_HEX = "0F2043", "1E407C", "B8943A"
LIGHT_HEX, GREEN_HEX, RED_HEX = "EEF2F8", "E8F5E9", "FDECEA"
AMBER_HEX, BLUE_HEX, ROW_HEX, WHITE_HEX = "FFF8E7", "E8F1FB", "F7F9FC", "FFFFFF"


def set_run_font(run, size=11, bold=False, color=DARK, italic=False, font=KR_FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/>'
            f"</w:tcMar>"
        )
    )


def set_table_borders(table, color="D0D7E2", sz="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    tbl_pr.append(
        parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f"</w:tblBorders>"
        )
    )


def set_left_accent(cell, color=NAVY_HEX, sz="24"):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="nil"/><w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="nil"/><w:right w:val="nil"/>'
            f"</w:tcBorders>"
        )
    )


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def cell_text(cell, text, size=10, bold=False, color=DARK, align="left"):
    cell.text = ""
    align_enum = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = align_enum
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1 if i < len(lines) - 1 else 0)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(line)
        set_run_font(run, size=size, bold=bold, color=color)
    set_cell_margins(cell)


def add_runs(paragraph, parts, size=11, color=DARK):
    for part in parts:
        if isinstance(part, str):
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)
        else:
            text, bold, *rest = part
            c = rest[0] if rest else color
            run = paragraph.add_run(text)
            set_run_font(run, size=size, bold=bold, color=c)


class Notes:
    def __init__(self):
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width = Mm(210)
        sec.page_height = Mm(297)
        sec.left_margin = Mm(16)
        sec.right_margin = Mm(16)
        sec.top_margin = Mm(16)
        sec.bottom_margin = Mm(16)
        sec.header_distance = Mm(8)
        sec.footer_distance = Mm(8)

        normal = self.doc.styles["Normal"]
        normal.font.name = KR_FONT
        normal.font.size = Pt(11)
        normal.font.color.rgb = DARK
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.18

        header = sec.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("9/17  AI·반도체 시장 코멘트  ·  FOMC 이후  ·  강의노트")
        set_run_font(r, size=8.5, color=GRAY)

        footer = sec.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("퀵코멘트+첨부 PDF를 주제별로 재구성  ·  숫자는 공개 자료로 교차검증  ·  ")
        set_run_font(r, size=8, color=GRAY)
        fld = parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE ">'
            f'<w:r><w:rPr><w:sz w:val="16"/><w:color w:val="4B5563"/>'
            f'<w:rFonts w:ascii="{KR_FONT}" w:hAnsi="{KR_FONT}" w:eastAsia="{KR_FONT}"/></w:rPr>'
            f"<w:t></w:t></w:r></w:fldSimple>"
        )
        fp._p.append(fld)

        core = self.doc.core_properties
        core.title = "9월 17일 AI·반도체 시장 코멘트"
        core.author = "준혁"
        core.subject = "FOMC, OpenAI DevDay, 하이닉스-인텔, 데이터센터 전력, 속도조절론"

    def p(self, text, size=11, bold=False, color=DARK, space_after=6, space_before=0, align="left"):
        para = self.doc.add_paragraph()
        para.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(space_before)
        para.paragraph_format.line_spacing = 1.18
        run = para.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
        return para

    def rich(self, parts, size=11, space_after=6, space_before=0):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(space_before)
        para.paragraph_format.line_spacing = 1.18
        add_runs(para, parts, size=size)
        return para

    def h1(self, text, num=None):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(14)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.line_spacing = 1.1
        if num:
            run = para.add_run(f"{num}  ")
            set_run_font(run, size=16, bold=True, color=GOLD)
        run = para.add_run(text)
        set_run_font(run, size=16, bold=True, color=NAVY)
        para._p.get_or_add_pPr().append(
            parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="12" w:space="4" w:color="{NAVY_HEX}"/>'
                f"</w:pBdr>"
            )
        )
        return para

    def h2(self, text):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(4)
        run = para.add_run(text)
        set_run_font(run, size=13, bold=True, color=NAVY2)
        return para

    def h3(self, text):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(8)
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(text)
        set_run_font(run, size=11.5, bold=True, color=NAVY)
        return para

    def bullet(self, text, level=0, bold_lead=None, size=11):
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.55 + level * 0.45)
        para.paragraph_format.first_line_indent = Cm(-0.35)
        para.paragraph_format.space_after = Pt(2.5)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing = 1.15
        mark = "• " if level == 0 else "– "
        run = para.add_run(mark)
        set_run_font(run, size=size, color=NAVY2 if level == 0 else GRAY)
        if bold_lead:
            run = para.add_run(bold_lead)
            set_run_font(run, size=size, bold=True, color=DARK)
            run = para.add_run(text)
            set_run_font(run, size=size, color=DARK)
        else:
            run = para.add_run(text)
            set_run_font(run, size=size, color=DARK)
        return para

    def flow(self, items, size=11):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.line_spacing = 1.2
        for i, item in enumerate(items):
            if i:
                run = para.add_run("   →   ")
                set_run_font(run, size=size, bold=True, color=GOLD)
            run = para.add_run(item)
            set_run_font(run, size=size, bold=True, color=NAVY)
        return para

    def callout(self, title, body, kind="key"):
        palette = {
            "key": (NAVY_HEX, LIGHT_HEX, NAVY),
            "bull": ("166534", GREEN_HEX, GREEN),
            "bear": ("991B1B", RED_HEX, RED),
            "note": (GOLD_HEX, AMBER_HEX, AMBER),
            "blue": (NAVY2_HEX, BLUE_HEX, NAVY2),
        }
        accent, fill, title_color = palette[kind]
        table = self.doc.add_table(rows=1, cols=1)
        table.autofit = True
        table.allow_autofit = True
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        shade_cell(cell, fill)
        set_left_accent(cell, accent, sz="28")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.space_before = Pt(0)
        r = p1.add_run(title)
        set_run_font(r, size=10, bold=True, color=title_color)
        if isinstance(body, str):
            body = [body]
        for line in body:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(line)
            set_run_font(r, size=10.5, color=DARK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def table(self, headers, rows, col_widths=None, first_col_bold=True):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table, color="D5DCE6", sz="4")
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            shade_cell(cell, NAVY_HEX)
            cell_text(cell, h, size=9.5, bold=True, color=WHITE, align="center")
        prevent_row_split(table.rows[0])
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = table.rows[r_i + 1].cells[c_i]
                shade_cell(cell, ROW_HEX if r_i % 2 else WHITE_HEX)
                align = "left" if c_i == 0 else "center"
                bold = first_col_bold and c_i == 0
                cell_text(cell, str(val), size=9.5, bold=bold, color=DARK, align=align)
            prevent_row_split(table.rows[r_i + 1])
        if col_widths:
            for row in table.rows:
                for i, w in enumerate(col_widths):
                    row.cells[i].width = Cm(w)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)
        spacer.paragraph_format.space_before = Pt(2)
        return table

    def image(self, name, width=16.4):
        path = CHARTS / name
        if not path.exists():
            self.p(f"[차트 없음: {name}]", size=9, color=GRAY)
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width))

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))


def build():
    n = Notes()

    n.p("2026. 9. 17. 강의노트  ·  FOMC 발표 다음 날  ·  퀵코멘트 + 첨부 13개 PDF", size=10.5, color=GRAY, align="center", space_after=4)
    n.p("FOMC 이후", size=13, bold=True, color=GOLD, align="center", space_after=2)
    n.p("9월 17일  AI·반도체 시장 코멘트", size=22, bold=True, color=NAVY, align="center", space_after=2)
    n.p("금리 경로  ·  속도조절론  ·  DevDay  ·  하이닉스–인텔  ·  전력 TAM", size=12, bold=True, color=NAVY2, align="center", space_after=8)
    n.p("학습용 GPU에서 추론·에이전트 배포로 논리가 한 단계 넓어지는 주", size=11, color=GRAY, align="center", space_after=10)

    n.callout(
        "한 장 결론",
        [
            "이번 인상은 침체형 인상이 아니다. 강한 경기 + 높은 인플레 + AI 자본수요다. 주식은 적응의 동물이다.",
            "속도조절론은 ‘개발 각도’ 이야기다. 45도에서 35도로 낮춰도 7월까지 보던 수요 자체가 줄어든다고 보면 안 된다.",
            "DevDay(9/29)에서 볼 것은 더 똑똑한 모델이 아니라 Agent 사용량·API 토큰·기업 배포다.",
            "하이닉스–인텔 오하이오는 탐색이다. Base Die → EMIB → Ohio가 계약으로 넘어가는지를 단계로 본다.",
            "전력은 2026년 +31%, 2028년부터 그리드 갭. AEMA는 병목 완화이지 변압기 대체가 아니다.",
        ],
        kind="key",
    )

    n.h2("이 노트를 읽는 법")
    n.table(
        ["표시", "뜻"],
        [
            ["사실", "연준 SEP, OpenAI, Reuters, TrendForce, 거래소 종가로 확인된 것"],
            ["부분 / 주의", "방향은 맞지만 숫자·시제·범위가 다른 것"],
            ["해석", "논리로 쓸 수는 있으나 확정처럼 쓰면 안 되는 것"],
            ["쓰지 말 것", "Ohio 메모리 확정, NVIDIA 12단→8단 전면 전환, AEMA=전력기기 수요 소멸, 속도조절=캡엑스 종료"],
        ],
        col_widths=[3.6, 14.0],
    )

    n.h2("오프닝 멘트 (녹화용)")
    n.p("오늘은 FOMC 다음날입니다. 금리 인상 자체는 선반영됐고, 시장이 받은 힌트는 ‘27년까지 기다릴 수 있다’입니다. 그와 동시에 하이닉스–인텔 오하이오, OpenAI DevDay, 데이터센터 전력 TAM이 한꺼번에 겹쳤습니다. 한 줄로 말하면 매크로는 부담, AI 수요의 무게중심은 학습에서 추론·에이전트 배포로 이동합니다.")

    n.h1("오늘 숫자 한 장", num="0.")
    n.table(
        ["항목", "숫자", "한 줄"],
        [
            ["FOMC", "3.75–4.00%  (+25bp)\n중앙값 4.1 / 4.1 / 3.9", "올해 한 번 더, 내년 동결이 중앙값"],
            ["GDP / Core PCE", "26년 2.3 → 27년 2.4\nCore PCE 27년 2.5 유지", "성장은 올리고 물가 경로는 유지"],
            ["9/16 美", "WTI 105.83 / 10Y 5.04%\nS&P −0.45 · SOX +0.4", "최악 매크로에도 반도체 차별화"],
            ["9/16 韓", "KOSPI 6,717.97 (+1.37%)\n하이닉스 +4.08% / 삼전 +2.01%", "전약후강, 반도체가 지수 지지"],
            ["DevDay", "9/29 샌프란시스코\n키노트 10:00 PT 라이브", "모델보다 Agent·API 배포"],
            ["Ohio", "임차 또는 CSP JV 탐색\n제품·계약 미정", "뉴스 ≠ 수주. 단계 확인"],
            ["DC 전력", "26년 161GW (+31%)\n30년 수요 490.7 vs 그리드 222.6", "갭은 종이 숫자. BTM·유연부하 감안"],
        ],
        col_widths=[3.2, 6.4, 8.0],
    )

    # ── 1. FOMC ──
    n.h1("FOMC — 인상은 선반영, 핵심은 경로", num="1.")
    n.callout(
        "한 줄 결론",
        [
            "「경기침체형 금리인상」이 아니라 「강한 경기 + 높은 인플레이션 + AI 투자에 따른 자본수요」다.",
            "중앙값은 올해 한 차례 추가 후 2027년 동결. 컨센 3~4회(27년까지)보다 낮은 경로다. 다만 점도표 분포는 중앙값보다 매파다.",
        ],
        kind="key",
    )
    n.h2("1) 공개 숫자 — 연준 SEP 9/16")
    n.table(
        ["항목", "2026", "2027", "2028", "장기", "판정"],
        [
            ["연방기금 중앙값", "4.1", "4.1", "3.9", "3.2", "사실"],
            ["6월 중앙값", "3.8", "3.6", "3.4", "3.1", "사실"],
            ["실질 GDP", "2.3", "2.4", "2.2", "2.0", "사실"],
            ["6월 GDP", "2.2", "2.3", "2.2", "2.0", "사실 · +0.1%p"],
            ["Core PCE", "3.4", "2.5", "2.2", "—", "27년 2.5 유지"],
            ["실업률", "4.1", "4.1", "4.1", "4.2", "사실"],
        ],
        col_widths=[3.4, 2.6, 2.6, 2.6, 2.4, 4.0],
    )
    n.image("01_fomc_dots.png")
    n.bullet("만장일치 +25bp, 목표 범위 3.75–4.00%. Warsh 의장은 자신의 점을 다시 제출하지 않았다. “forward guidance business가 아니다.”")
    n.bullet("2026년 점: 4.125 부근 12명, 4.375 부근 4명, 3.875 부근 2명. 추가 1회가 중앙값, 추가 2회도 4명.")
    n.bullet("2027년 점: 4.375 부근이 8명으로 늘어난다. ‘내년 동결’은 중앙값 이야기이지 전원 합의가 아니다.")

    n.h2("2) 코멘트에서 고친 것")
    n.table(
        ["코멘트", "판정", "메모"],
        [
            ["27년 성장 2.4 > 26년 2.3", "사실", "Q4/Q4 중앙값. Warsh 기자회견과 동일"],
            ["Core PCE 결국 2.5, 이전과 동일", "사실", "2027 중앙값 6월·9월 모두 2.5"],
            ["올해 추가 1회 후 내년 동결", "중앙값 사실", "상단은 4.375. 4.5% 점수는 없음"],
            ["장기 3.0% 부근", "부분", "장기 중앙값 3.2. 3.0에 찍은 위원은 6명"],
            ["컨센 3~4회보다 낮다", "해석", "시장 컨센과 점도표를 비교한 읽기"],
            ["27년 하반기 인하 여지", "해석", "점도표 자체는 27년 4.1 유지"],
        ],
        col_widths=[5.2, 2.8, 9.6],
    )

    n.h2("3) 어떻게 읽을까")
    n.bullet("2022년과 구별한다. 그때는 전방위 물가에 해소 경로가 안 보였다. 지금은 원자재가 서비스로 번지는 것을 막으려는 선제 인상이다.", bold_lead="2022 ≠ 지금  ")
    n.bullet("역사적으로 금리 인상기에 주가가 더 오른 경우가 많다. 기업 실적이 받치면 그렇다. 22년은 전쟁+메모리 공급초과가 겹친 예외에 가깝다.", bold_lead="적응  ")
    n.bullet("Higher for Longer 헤드라인보다, 성장이 버티고 물가 경로가 유지되면 ‘적응하고 나면 기다려볼 만하다’로 심리가 바뀔 수 있다.", bold_lead="심리  ")
    n.bullet("금리 상단이 보이면 스마트머니는 미리 움직인다. 지금 당장 성장주 프리미엄이 돌아온다는 뜻은 아니다.", bold_lead="상단  ")
    n.bullet("지정학이 다시 커지지 않는다면, 본게임은 AI 생산성 효과와 연준 AI 태스크포스(연말 검토)다.", bold_lead="본게임  ")
    n.callout("한 줄", ["인상 여부보다 추가 경로. 중앙값은 기다릴 수 있다. 분포는 더 매파다. 유가·10년물이 다시 세게 가면 이 읽기는 흔들린다."], kind="blue")

    # ── 2. 시장 ──
    n.h1("9/16 시장 — 매크로 부담, 반도체 차별화", num="2.")
    n.image("04_macro_vs_chips.png")
    n.h2("1) 미국 — 새벽 장")
    n.table(
        ["지표", "숫자", "의미"],
        [
            ["WTI / Brent", "$105.83 (+4.38%) / $108.75", "사우디 송유관·얀부 차질. 브렌트 5/19 이후 최고"],
            ["美 10년", "장중 5.041%", "2007 이후 약 19년래 최고. 레벨보다 속도"],
            ["Dow / S&P / Nasdaq", "−0.63 / −0.45 / −0.78", "금리·유가 부담"],
            ["SOX", "약 +0.4%", "전일 급락 후 반등. 코멘트 +0.5~0.6는 약간 큼"],
            ["개별", "MU +0.4 · COHR +1.8 · DELL +1.7", "CapEx 수혜는 반발 매수"],
            ["크립토", "COIN −10.1 · CRCL −11.4", "CLARITY 절차표결 부결"],
        ],
        col_widths=[3.6, 6.4, 7.6],
    )
    n.callout(
        "투자 관점",
        [
            "Macro: 유가↑ → CPI 기대↑ → 금리↑ → PER 부담",
            "AI: 투자 지속 → GPU/ASIC → HBM·DRAM",
            "시장은 ‘AI 수요 둔화’보다 ‘고금리 밸류에이션 조정’을 가격에 넣고 있다.",
        ],
        kind="blue",
    )

    n.h2("2) 웰스파고 — AI 수요가 아니라 쏠림")
    n.bullet("S&P 연말 목표 7,950 → 7,700. 주식 비중 72%(1969 이후 최고 추정), 10년 5%대.")
    n.bullet("도달 전 5–10% 조정 가능. 기술 동일비중, 헬스케어 확대. 중간선거 후 DC 규제 리스크.")
    n.bullet("적정 주식 60 / 채권 40이라는 설명. AI·반도체 수요 둔화가 하향 근거는 아니다. 웰스파고 전망일 뿐이다.")

    n.h2("3) 한국 — 전약후강, 6,700 지지")
    n.table(
        ["항목", "숫자", "판정"],
        [
            ["KOSPI / KOSDAQ", "6,717.97 (+1.37%) / 815.98 (+0.44%)", "사실 (거래소)"],
            ["4일 연속 하락 후 반등", "기관 +1.21조, 외인 −1.68조", "사실"],
            ["SK하이닉스 / 삼성전자", "+4.08% / +2.01%", "사실. 코멘트 +3.6/+1.8은 약간 작음"],
            ["카카오페이", "약 −11.5%", "사실. CLARITY 연동"],
            ["피에스케이 +9.5%, 효성중공업 +4.4%, LS에코 +13.5%", "코멘트", "테마 설명용. 종가는 방송 전 재확인"],
        ],
        col_widths=[5.6, 6.4, 5.6],
    )
    n.flow(["유가 $105+ · 10년 5%", "Risk-off로 안 번짐", "반도체 중심 반등", "KOSPI 6,700 지지"])
    n.bullet("촉매는 하이닉스–인텔 협력 논의. 소부장·전력·케이블·광통신이 같이 돌았다.")
    n.bullet("대미투자 국회 보고 연기는 두산에너빌리티·한전기술에 부담으로 읽혔다.")
    n.callout("핵심 구도", ["금리·유가 부담은 남는다. 그래도 반도체·AI 인프라 수급이 지수를 받쳤다. FOMC 이후에는 9월 인상 여부보다 추가 경로."], kind="key")

    # ── 3. 속도조절 ──
    n.h1("속도조절론 — 각도는 낮춰도 수요는 안 줄어든다", num="3.")
    n.callout(
        "한 줄 결론",
        [
            "속도조절 얘기 자체가, 속도가 너무 빠르기 때문에 나온 것이다.",
            "프론티어 학습을 숨 고르기 해도 추론·에이전트 수요가 더 크고 더 빠르다. 가드레일 보강 자체가 연산 수요를 늘린다.",
        ],
        kind="key",
    )
    n.h2("1) 35도 → 45도 → 다시 30~35도")
    n.p("4–6월, Astra 학습 중일 때 시장이 보던 개발 각도를 35도, 수요를 100이라고 하자. 9월 Astra 공개 후 각도가 45도 이상으로 보였고 안전 이슈가 붙었다.")
    n.flow(["35도 (4–6월 예상)", "45도+ (Astra 충격)", "30~35도 (속도조절)", "수요 100은 유지"])
    n.bullet("프론티어(Astra급 이상)는 비싸다. 업계가 숨 고르기 하고, 돈이 되는 추론 토큰에 집중하는 것이 모델사 입장에서 베스트 전략으로 읽힌다.")
    n.bullet("중국이 바로 못 따라온다고 보면 신사협정·가드레일의 유인도 있다. 다만 중국이 같이 멈춰 줄지는 별개다. 그래서 당위론만으로 속도조절을 현실 정책으로 단정하면 안 된다.")
    n.bullet("이미 만들어 둔 프론티어만으로도 수요가 넘친다. 무게중심은 학습에서 추론으로 이동 중이다.")

    n.h2("2) Amodei를 상상해 적은 메모")
    n.bullet("속도가 예상을 크게 넘는다. 경쟁사와 같이 숨 고르고 싶다.")
    n.bullet("중국은 딜레마다. 가드레일 합의에는 시간이 걸린다. 그동안 본인·경쟁사(Grok, 구글, 중국)가 진짜로 개발을 늦출지는 자신 없다.")
    n.bullet("합의가 돼도 추론은 따로 는다. 에이전트 때문에 추론이 전체의 절반 이상, 더 빠르다.")
    n.bullet("논란 속에 반도체 가격이 하향되면 원가 부담이 줄어 모델사에는 오히려 좋다 — 이건 상상 메모다.")

    n.h2("3) 핵심 문장")
    n.callout(
        "Amodei 쪽 논리에서 가져갈 두 문장",
        [
            "AI 발전이 이토록 빠를 것이라고는 예상하지 못했다.",
            "지금 상태에서 동결해도, 현재 활용 가치는 전체의 5~10%에 불과하다.",
        ],
        kind="note",
    )
    n.p("즉 속도조절은 ‘AI를 그만두자’가 아니다. 이미 있는 모델의 상용화·추론이 본게임이라는 뜻이다.")

    n.h2("4) CapEx 딜레마")
    n.flow(["모델 경쟁", "학습비 ↑", "GPU·메모리", "DC CapEx", "전력", "실적·멀티플"])
    n.bullet("비용 증가 속도 > 수익화 속도면 투자 지속성 → CapEx 증가율 → 인프라 수요 기대가 한꺼번에 내려간다.")
    n.bullet("반론: 생산성·신규 매출이 비용을 정당화하거나, 칩 성능↑·전력효율↑·단가↓가 동시에 오면 완화된다.")
    n.bullet("규제 강화 ≠ CapEx 둔화. 국가 전략산업으로 투자가 이어질 수 있다.")
    n.bullet("Politico(수요일): 일시 중단 50% vs 계속 발전 31%. 여론은 가속보다 억제. 여론 ≠ 캡엑스 결정.")
    n.callout("Broadcom Hock Tan", ["2027–28 AI 반도체 매출 전망을 수정할 이유 없다. 속도조절론 ≠ 인프라 투자 둔화. 학습보다 추론 상용화를 낙관. Anthropic이 27–28년 최대 커스텀 칩 고객. 사이클 Peak-out보다 Duration Extension."], kind="bull")

    # ── 4. DevDay ──
    n.h1("OpenAI DevDay — 모델보다 Agent 배포", num="4.")
    n.callout(
        "한 줄 결론",
        [
            "DevDay 2026은 9월 29일, 샌프란시스코 Fort Mason. 공식 페이지 기준 개발자·API·도구의 실제 배포 행사다.",
            "볼 것은 ‘또 하나의 더 똑똑한 모델’이 아니라 Agent 사용량, API/토큰, 기업 배포 속도다.",
        ],
        kind="key",
    )
    n.h2("1) 사실")
    n.bullet("화 9/29, 키노트 10:00 PT(Sam Altman), 라이브 무료. 대면 등록 $650, 신청은 이미 마감.")
    n.bullet("공식 문구: technical sessions, test what’s new, OpenAI tools. 모델 발표 행사라고 단정하면 안 된다.")
    n.h2("2) 반도체로 읽는 3가지 — 해석")
    n.table(
        ["축", "내용", "인프라"],
        [
            ["① Agent 확대", "답변 → 웹·코딩·업무 수행", "토큰·추론량 증가"],
            ["② API 생태계", "기업이 자기 서비스에 모델 탑재", "AI 서버 수요"],
            ["③ 대규모 Deployment", "실험 → 기업 업무 배포", "GPU·DRAM·HBM 지속"],
        ],
        col_widths=[3.6, 7.4, 6.6],
    )
    n.flow(["Astra 광범위 배포", "DevDay", "Agent 대량 깔림", "추론 컴퓨팅 + 메모리"])
    n.bullet("코멘트: Agents API, 금융용 ChatGPT, GPT-Live를 잇달아 냈다. 공개 일정은 방송 전 재확인.")
    n.bullet("Astra가 가장 강한 광범위 배포 모델이며 사이버보안에서 처음 Critical로 평가됐다는 것은 코멘트. 확정 사실처럼 쓰지 말 것.")
    n.h2("3) DevDay에서 확인할 세 숫자")
    n.bullet("Agent 사용량이 얼마나 늘었는가.")
    n.bullet("API 호출·토큰이 얼마나 늘었는가.")
    n.bullet("기업용 Agent가 얼마나 빨리 배포되는가.")
    n.callout("CapEx 논리의 확장", ["이 세 개가 확인되면 ‘학습용 GPU 투자’에서 ‘추론용 컴퓨팅 + 메모리 투자’로 한 단계 더 넓어진다."], kind="blue")

    n.h2("4) 평가 기준이 바뀐다 — IQ → 실전")
    n.bullet("벤치마크 평균 점수보다 최종 성공률. 단계 성공 95%여도 20턴이면 0.95²⁰ ≈ 36%.")
    n.bullet("그래서 중간 실패를 재시도하며 고치는 능력이 핵심이다. Agentic AI → 추론·메모리·컴퓨팅 지속.")
    n.bullet("평가 변경 후 미·중 격차 약 7주 → 약 3개월. 비공개·장기 업무일수록 활용 능력 차이가 난다. (코멘트)")
    n.bullet("미국은 성능+안전, 중국은 성능 집중이라는 분석. 미국 AI 우위는 달러·미국 자산의 구조 경쟁력으로도 읽힌다 — 해석.")
    n.h3("애플도 같은 축")
    n.bullet("M8 Ultra 2~4개 기업용 추론 서버 검토, NVLink Fusion 결합 가능성. 초기 검토, 2029 이후 옵션.")
    n.bullet("새 Siri는 온디바이스 실행형 비서. Instinct·Meta Muse 같은 웹 행동형 비서와 경쟁. 프라이버시가 제약일 수 있다. (WSJ)")

    # ── 5. 공포 ──
    n.h1("AI 공포 — 생각 다 같이 해야, 맡기지 말고", num="5.")
    n.callout(
        "한 줄 결론",
        [
            "과장과 자기이익(대중 경쟁 우위, 추론 집중)이 핵심 배경일 수 있다.",
            "인류애 담론보다 ‘소수 독점 결정 방식’이 더 큰 이슈다. FACT는 컴퓨팅이 턱없이 부족하다는 것이다.",
        ],
        kind="key",
    )
    n.bullet("가드레일은 필요하다. 다만 만드는 과정에 다양한 목소리가 들어가야 하고, 그러면 느려진다.")
    n.bullet("종말 시나리오의 다급한 주장은 규제를 서두르는 쪽에 유리하다. 가상 위험에 에너지를 쓰면 현재 피해를 놓친다.")
    n.bullet("가까운 미래 통제 불능 AI의 실존 위험은 낮다는 쪽 코멘트. 일부는 가능성 자체에 의문을 단다.")
    n.bullet("연구자 사임·10년 내 멸종 10%+ 발언은 속도조절론의 재료가 됐다. 투자로 번역할 것은 ①멸종 10%가 아니라 ②정치·규제화 → ③개발 속도 → ④그래도 추론 수요다.")

    n.h2("술레이만 vs Anthropic — 의식·권리")
    n.table(
        ["쪽", "입장", "함의"],
        [
            ["Suleyman / Nadella", "의식·권리·감정을 학습시키지 말 것. research, focus, deliberate pacing", "도구로 유지"],
            ["Anthropic 가이드", "의식 불확실 → moral patient·model welfare·rights 논의 가능", "가능성을 연다"],
            ["Zuckerberg / Jensen", "조직적으로 늦추는 접근에 부정", "속도 유지"],
        ],
        col_widths=[4.2, 8.0, 5.4],
    )
    n.bullet("과학적 확정이 아니라 안전 설계 철학 차이다. ‘의식이 있는가’와 ‘모르는 상태에서 어떻게 훈련할 것인가’는 다른 질문이다.")
    n.callout("한 줄", ["공포를 빅테크에 맡기지 말 것. 규칙 결정권을 소수가 가져가면, 그 자체가 투자·규제 리스크다."], kind="note")

    # ── 6. Hynix-Intel ──
    n.h1("하이닉스–인텔 — Ohio는 탐색, Base Die는 별축", num="6.")
    n.callout(
        "한 줄 결론",
        [
            "Reuters 9/16: 미국 첫 메모리 생산을 인텔과 탐색. ①Ohio 일부 임차 ②인텔+대형 CSP JV.",
            "하이닉스 ‘결정된 사항 없다’, 인텔 ‘추측’. 산자부: 국가핵심기술이면 산업기술보호법 심사.",
            "오늘 뉴스로 소부장을 바로 연결하지 말 것. Base Die → EMIB → Ohio 계약을 단계로 본다.",
        ],
        kind="key",
    )
    n.h2("1) 사실과 경계")
    n.bullet("Ohio One은 2022년 최대 $1000억 구상. 두 공장 완공은 2030–31로 지연된 상태.")
    n.bullet("하이닉스는 인디애나에 $40억+ HBM 첨단 패키징. 2029 하반기 차세대 HBM 목표. 후공정이다.")
    n.bullet("한국 $3500억 대미 약속 중 조선 $1500억, 나머지 $2000억은 협상 중 — Reuters 맥락.")
    n.bullet("생산 제품(DRAM인지 HBM인지)은 Reuters도 확인 못 했다.")

    n.h2("2) 네 시나리오 — 아직 메뉴판")
    n.table(
        ["시나리오", "하는 일", "신규 수요", "지금"],
        [
            ["A. Base Die 파운드리", "로직 Base Die", "선단 로직 장비", "업계 검토. Ohio 기사와 별개"],
            ["B. DRAM 전공정", "하이닉스 DRAM 웨이퍼", "DRAM 전공정 장비 대규모", "가장 파급 큼. 미확정"],
            ["C. HBM 일부/전체", "DRAM→HBM", "DRAM+HBM 후공정", "미확정"],
            ["D. 단순 패키징", "HBM 패키징", "본딩·검사", "인디애나가 이미 이 축"],
        ],
        col_widths=[4.0, 4.4, 5.2, 4.0],
    )
    n.flow(["한국 DRAM 전공정", "Intel Ohio 전공정?", "Indiana HBM 패키징", "Intel Foundry Base Die?"])

    n.h2("3) 왜 Base Die가 새로운 수혜 영역인가")
    n.p("HBM4: 하이닉스 DRAM Die + TSMC Base Die. HBM4E: TSMC + Intel Foundry 가능성이 거론된다. 목적은 TSMC 의존 완화.")
    n.bullet("Intel Foundry가 진짜로 Base Die를 찍으면 웨이퍼→증착→식각→세정→검사→패키징이 새로 열린다.")
    n.bullet("국내에서 인텔 거래선이거나 미국 Fab 공급 경험이 있는 업체를 따로 볼 가치가 있다. 예: 이오테크닉스, 코미코, 인텍플러스, 피에스케이, 주성엔지니어링.")
    n.bullet("장기 옵션은 EMIB. GPU(TSMC) + HBM4E(하이닉스) + Intel EMIB. 채택돼야 2차 이야기다.")
    n.callout(
        "TSMC에게 제일 아픈 그림 (아직 시나리오)",
        ["하이닉스 HBM 스택 + Intel 18A Base Die + Intel EMIB + GPU/ASIC. 인텔이 하청이 아니라 Base Die+패키징+HBM 통합을 가져간다."],
        kind="note",
    )

    n.h2("4) HBM 디스펙 ≠ 수요 감소")
    n.p("시장 오해: 단수 ↓ = HBM 수요 ↓. 코멘트 구조: 역할 재정의 + 메모리 계층화.")
    n.table(
        ["계층", "역할"],
        [
            ["GPU ↔ HBM", "초고대역폭·저지연. 용량은 줄여도 대역폭은 유지·상향"],
            ["DDR / SOCAMM", "대용량 메모리"],
            ["CXL / CMX", "확장·재사용 데이터"],
            ["SSD", "장기 저장 / 대규모 KV Cache"],
        ],
        col_widths=[4.4, 13.2],
    )
    n.bullet("8단 옵션 확대는 확인 가능(삼성 HBM4E 12단 48GB·8단 32GB·16단 64GB 라인업). NVIDIA 12→8 전면 전환은 아직 관측.")
    n.bullet("낮은 적층 → 공급량 ↑ → GPU 출하 ↑ → 스택 총량·Base Die 총수요 ↑ 라는 2차 효과도 가능하다.")
    n.bullet("Base Die가 중요해질수록 가치가 메모리 업체가 아니라 NVIDIA·Broadcom·Marvell 쪽 컨트롤러/IP로 갈 수 있다. ‘삼성 유리’는 ④⑤가 약하다.")

    n.h2("5) 확인 체크리스트")
    n.bullet("Ohio 제품이 DRAM인가, 패키징인가, 미정인가.")
    n.bullet("HBM4E Base Die의 Intel Foundry 물량이 계약으로 나오는가.")
    n.bullet("산자부 국가핵심기술 심사 대상인가.")
    n.bullet("인디애나 패키징 일정(29H2)이 유지되는가.")

    # ── 7. Power ──
    n.h1("데이터센터 전력 — TAM과 유연 부하", num="7.")
    n.callout(
        "한 줄 결론",
        [
            "TrendForce: 2026년 161GW(+31%), AI 서버 33.4%. 2027도 +30%대. 2028부터 그리드 갭이 본격화.",
            "2030 수요 490.7 vs 그리드 222.6, 종이 갭 268GW(미국 170GW+). 현장발전은 분모에 없다.",
            "AEMA는 고정 부하를 유연 부하로 바꿔 접속을 앞당긴다. 변압기·전선을 대체하지 않는다.",
        ],
        kind="key",
    )
    n.image("02_dc_power_gap.png")
    n.image("03_ai_server_share.png")
    n.h2("1) TrendForce 숫자")
    n.table(
        ["연도", "수요 용량", "AI 서버 비중", "일반 서버 비중"],
        [
            ["2025", "122.9GW", "약 25%", "약 40%"],
            ["2026", "161GW (+31%)", "33.4%", "하락 구간"],
            ["2027", "+30%대 유지", "40% 넘을 가능성", "25.6%"],
            ["2030", "490.7GW", "—", "—"],
        ],
        col_widths=[3.2, 4.6, 5.0, 4.8],
    )
    n.bullet("4분류: 일반 서버 / AI 서버 / 기타 IT / 냉각·전력 등 비IT. 효율은 좋아지지만 AI 서버+비IT가 같이 는다. HVDC 검토 시작.")
    n.bullet("2025까지 그리드는 수용. 2026부터 어긋나고 2028 이후 확대. 미국은 28년까지 상당수가 그리드에 붙고, 이후 지연이 잦아진다.")
    n.bullet("종이 갭 > 실제 부족일 수 있다(BTM 미포함). 반대로 계통·송배전 지연은 실제 부족을 만든다.")

    n.h2("2) AEMA — 9/16 출범")
    n.bullet("사실: NVIDIA · Google · Emerald AI. 론칭 파트너에 Anthropic, National Grid, AES, RWE, Constellation, NRG 등.")
    n.bullet("기존: 500MW면 항상 500MW. AEMA: 평소 500, 부족 시 350–400으로 잠시 줄였다가 복귀.")
    n.bullet("수단: 워크로드 이동·일시 정지, 배터리, 현장 발전. 기술 중립, 성능(속도·지속·예측·비상)으로 평가.")
    n.bullet("100GW는 AEMA/NVIDIA 주장. 2025말 유틸리티 발전 1,281GW의 약 7.8%라는 환산은 산수로는 맞다.")
    n.bullet("Emerald Conductor 256 GPU에서 25%·3시간, 최근 1분 내 40%는 코멘트/회사 실증. 공식 블로그에는 그 숫자가 없다.")
    n.flow(["고정 100% 부하", "수년 접속", "유연 70–80%", "접속 단축 + 기존 망 활용"])
    n.callout("투자로 번역", ["NVIDIA가 GPU에서 AI Factory + 전력 운영 플랫폼으로 영역을 넓히는 흐름. 절대 전력·변압기 수요는 DC가 늘어면 같이 는다. AEMA는 병목을 깎는 기술이다."], kind="blue")

    n.h2("3) 리드타임 — 병목은 칩이 아니다")
    n.table(
        ["부품", "상태", "지금 / 균형"],
        [
            ["ABF 기판", "Very Tight", "48–56주 / 12주"],
            ["HDD", "Very Tight", "50주 / 16주"],
            ["DRAM", "Very Tight", "20주 / 8주"],
            ["MLCC", "Tight", "30주 / 12주"],
            ["NAND eSSD", "Tight", "16주 / 8주"],
            ["GPU", "Balanced", "20–30주 / 20–30주"],
        ],
        col_widths=[4.0, 4.0, 9.6],
    )
    n.p("TrendForce 기준. GPU는 균형, 기판·HDD·메모리가 더 타이트하다.", size=10.5, color=GRAY)

    # ── 8. 개별 ──
    n.h1("개별 — 소부장 · 기판 · OSAT · 전력", num="8.")
    n.h2("1) 오늘 장에서 같이 본 것")
    n.bullet("삼성: Taylor Tesla AI 칩 시제품. cHBM Base Die 자체+TSMC 2-track. 수율 검증 단계.")
    n.bullet("두산: 3년 7,551억 고급 CCL. 가동률 100%+. AI 가속기·고속망 PCB 소재.")
    n.bullet("삼성SDS: Anthropic Select Tier Claude 국내 최초. 기업용 구축.")
    n.bullet("마이크론: 세계 첫 512GB DDR5 RDIMM 9,200MT/s, 전력 60%↓, AMD·인텔 검증, 27H2 양산. Sadana: 메모리가 AI 상한, 실질 신규 capa는 2028 이후. FY27 캡엑스 $450억+.")
    n.bullet("Dell CFO: 대형 고객 주문이 수년 치까지 논의. 장기 AI 서버 백로그.")

    n.h2("2) 대덕전자 LTA")
    n.bullet("Tesla · Marvell · Renesas 협의. Renesas는 문구 막판, Tesla는 선수금 큰 틀.")
    n.bullet("이미 확정 1곳(Amkor 추정). Marvell은 연장+자금 조건. 선수금 총액이 올해 공시 8,000억 투자를 웃돈다는 업계 전언.")
    n.bullet("질적 변화: 단순 공급 → 선수금 기반 선제 CAPA. 다만 기판 LTA는 Take or Pay보다 약하다. 매출 전환이 후속 확인.")

    n.h2("3) 하나마이크론")
    n.bullet("삼성 DDR5 외주 + 하이닉스 베트남 외주 → 가동률. 4Q26부터 베트남 증설. FY26 1조+ → FY27 1.2조.")
    n.bullet("비메모리 2Q26 QoQ +40%, 27년 Line 5. 12M fwd P/E 10배 이하 vs 목표 10배 후반 — 증권가. SFA와 키 맞추기 성격도.")

    n.h2("4) 조선 — 이익↑ 멀티플↓")
    n.bullet("3사 합산 OP 1Q26 1.6조 → 2Q26 2.1조(+30%). 시총 5/26 146조 → 9/16 89조(−39%).")
    n.bullet("12M fwd P/E 14배 미만(현대중 12.9, 한화오션 12.6, 삼성중 13.4). 23년 4월 이후 최저.")
    n.bullet("리레이팅 3축: 4행정 중속 엔진(AIDC) · FDC · 글로벌 함정. 상선 P·Q 한계와 별개로, 신규 축이 수주로 확인돼야 한다.")

    # ── 9. 포트 ──
    n.h1("포트폴리오 · 환율 · 정치", num="9.")
    n.bullet("7월 말 바닥은 지났다. 이 정도 악재에 바닥 대비 +20%면, 6,000 아래 쌍바닥은 새 악재가 필요하다.")
    n.bullet("조정은 전저점 회귀가 아니라 숨고르기. 수급이 얇아 거래량만으로도 흔들린다. 일일 변동에 과반응하지 말 것.")
    n.bullet("바벨 예: AI 40–50 / Non-AI(조선·유통·일부 금융) 30 / 현금성 20–30. 편한 비중.")
    n.bullet("Anthropic IPO 10–11월은 수급을 뺏을 수 있다. SpaceX(6/12) 때는 코스피 8,000+, 하이닉스 200만 초였다.")
    n.bullet("다음 큰 숫자: 10/1 마이크론 실적부터 실적 시즌. 중간선거는 11/3.")
    n.bullet("원/달러 1,300원대 초반이 바닥이었을 가능성. 4Q 대미투자 달러 수요+추가 인상으로 재상승 여력. 백화점 외인 매수. 원화 강세 배팅은 재검토.")
    n.bullet("Blue Wave ≠ 전국 DC 중단. 뉴욕은 중단, 메인·버지니아 민주 지사는 모라토리엄을 안 받았다. 현실 영향은 미허가 승인 장기화·비용.")
    n.bullet("CLARITY는 폐기보다 일정 지연. 하원 2025 통과, 상원 제동. SEC vs CFTC. 스테이블 틀은 2025 GENIUS와 별개.")

    n.h1("보론 — 올릭스 Corporate Day", num="10.")
    n.bullet("오버행 해소 + 하반기 임상·L/O. 핵심은 ALK7 비만(OLX501A, siRNA로 ALK7 억제).")
    n.bullet("4Q 경쟁사 데이터: Alnylam ALN-2232, Arrowhead ARO-ALK7(±마운자로), GSK/Siran SA030.")
    n.bullet("타깃 유효성이 증명되면 OLX501A L/O 가치에 영향. 오버행 해소 = 주가 상승은 신한 전망. 실제는 임상·L/O.")

    n.h1("클로징", num="11.")
    n.callout(
        "오늘 가져갈 네 문장",
        [
            "1) FOMC: 인상은 소화됐다. 중앙값은 올해 한 번 더·내년 동결. 성장은 올리고 Core PCE 2.5는 유지. 분포는 매파다.",
            "2) 속도조절: 각도를 낮춰도 7월까지 보던 수요는 안 줄어든다. 추론·에이전트가 더 크고, 가드레일도 연산을 늘린다.",
            "3) DevDay(9/29): Agent 사용량 · 토큰 · 기업 배포. 확인되면 캡엑스는 학습 GPU에서 추론 컴퓨팅+메모리로 확장.",
            "4) 하이닉스–인텔과 전력: Ohio는 탐색, Base Die는 별축, 161GW는 사실. 계약과 그리드 갭을 단계로 본다.",
        ],
        kind="key",
    )
    n.h2("클로징 멘트 (녹화용)")
    n.p("주식은 적응의 동물입니다. 예상된 인상, 예상된 발언이면 심리는 ‘27년 기다릴 수 있다’로 갑니다. 그 사이 AI 수요는 학습 각도 논쟁과 따로 움직입니다. DevDay에서 에이전트가 얼마나 깔리는지, Ohio와 Base Die가 계약이 되는지, 전력이 그리드와 BTM과 유연 부하로 어떻게 메워지는지. 이 세 개가 올해 가을의 확인 포인트입니다.")
    n.p("— 9월 17일 노트. 원문 퀵코멘트(9/16 23:37 ~ 9/17 17:56)와 첨부 PDF를 주제별로 재구성. 종가·SEP·공식 페이지로 숫자를 고쳤다.", size=9.5, color=GRAY, align="right")

    n.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({OUT_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
