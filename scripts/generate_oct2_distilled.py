#!/usr/bin/env python3
"""10월 2일 코멘트를 큰 흐름 한 편으로 다시 쓴다.

숫자는 oct2_audit.py가 다시 계산한 값만 쓴다.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import generate_oct2_lecture as lec
import oct2_audit
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt

OUT_PATH = Path("/workspace/lectures/10월 2일 큰 흐름.docx")
KR_FONT = lec.KR_FONT
NAVY = lec.NAVY
NAVY2 = lec.NAVY2
GOLD = lec.GOLD
GRAY = lec.GRAY
DARK = lec.DARK
NAVY_HEX = lec.NAVY_HEX
set_run_font = lec.set_run_font


class Notes(lec.Notes):
    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width = Mm(210)
        sec.page_height = Mm(297)
        sec.left_margin = Mm(16)
        sec.right_margin = Mm(16)
        sec.top_margin = Mm(16)
        sec.bottom_margin = Mm(16)
        sec.header_distance = Mm(8)
        sec.footer_distance = Mm(8)

        normal = self.doc.styles["Normal"]
        normal.font.name = KR_FONT
        normal.font.size = Pt(11)
        normal.font.color.rgb = DARK
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.18

        header = sec.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("10/2 큰 흐름  ·  금리 천장 · AI 체인 · 피크 이후")
        set_run_font(r, size=8.5, color=GRAY)

        footer = sec.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("10월 2일 큰 흐름  ·  숫자는 검증된 입력만  ·  ")
        set_run_font(r, size=8, color=GRAY)
        fld = parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE ">'
            f'<w:r><w:rPr><w:sz w:val="16"/><w:color w:val="4B5563"/>'
            f'<w:rFonts w:ascii="{KR_FONT}" w:hAnsi="{KR_FONT}" w:eastAsia="{KR_FONT}"/></w:rPr>'
            f"<w:t></w:t></w:r></w:fldSimple>"
        )
        fp._p.append(fld)

        core = self.doc.core_properties
        core.title = "10월 2일 큰 흐름"
        core.author = "준혁"
        core.subject = "금리 천장, AI 인프라 체인의 확산, 메모리 피크 이후"


def build():
    oct2_audit.build()
    n = Notes()

    n.p("2026. 10. 2. 종가  ·  큰 흐름", size=10.5, color=GRAY, align="center", space_after=4)
    n.p("금리 천장  ·  AI 체인  ·  피크 이후", size=13, bold=True, color=GOLD, align="center", space_after=2)
    n.p("10월 2일", size=22, bold=True, color=NAVY, align="center", space_after=8)

    n.callout(
        "한 문장",
        [
            "지수는 10년물 5%대가 막고, 돈은 GPU를 지나 메모리·기판·광·전력으로 퍼지고 있다.",
            "메모리 주가의 싸움은 피크가 오느냐가 아니라, 피크 이후 이익을 몇 배에 사느냐다.",
        ],
        kind="key",
    )
    n.flow(["금리 5%대", "AI 캡엑스 확산", "피크 이후 몇 배"])

    n.h1("세 층", num="1.")
    n.table(
        ["층", "지금 벌어지는 일", "포지션"],
        [
            ["천장", "고용은 둔화, 주식은 상승, 10년물은 5.28%", "지수 베타는 제한. 현금 20%"],
            ["체인", "GPU 다음이 HBM, 기판, 광, 전력", "AI 50~60%. 삼전·닉스가 핵"],
            ["가격", "메모리 100→122→90. 이익은 60조대에 남는가", "8배가 열려야 224만원"],
        ],
        col_widths=[2.2, 8.6, 6.8],
    )
    n.p("27년 3월까지는 이 세 층이 계단으로 움직인다고 보고, 10년물이 5% 아래로 추세적으로 내려올 때 폭을 키운다.")

    n.h1("천장은 금리다", num="2.")
    n.p("9월 고용 +2.9만은 침체 진입이 아니다. 본론은 임금이 연초 3.7%에서 3.0%로 내려온 것이고, 근원 PCE는 3.0~3.2%에 남아 있다. 그래서 10월은 동결 쪽이고 인하 재료는 아니다. 당일 SOX는 +2.4%, 10년물은 5.28%였다. 9월 11일부터 10월 2일까지 10년물이 약 30bp 오르는 동안에도 SOX는 +11%, 마이크론은 +10%였다.")
    n.callout(
        "이 층의 결론",
        [
            "지수 전체의 추가 상승은 10년물 5%대가 정한다.",
            "그 안에서 남는 알파는 AI 직간접이다. 신규는 분할이고, 유가 100달러 위에서는 그 분할을 유지한다.",
            "10년물이 5% 아래로 자리 잡으면 베타를 늘린다. 12월 인상 확률이 다시 살아나면 계단으로 돌아간다.",
        ],
        kind="blue",
    )

    n.h1("돈은 체인 뒤로 퍼진다", num="3.")
    n.flow(["GPU", "HBM · 메모리", "기판 · 패키징", "스위치", "Copper", "Optical", "CPO"])
    n.p("같은 캡엑스가 종목마다 다른 가격으로 도착한다. 앞단은 이익의 높이, 뒷단은 2028~2030년에 공장이 서는지, 옆길은 이미 올라간 배수가 문제다.")
    n.table(
        ["자리", "이름", "이 흐름에서 하는 일"],
        [
            ["앞단", "삼성전자 · 하이닉스 · 마이크론", "가격이 한 번 더 오른 뒤 조정. 주가는 28년 정상 이익의 배수"],
            ["기판", "삼성전기 · 심텍 · 대덕 · 코리아써키트 · 인텍플러스", "신코·유니마이크론이 2028~2030 캐파를 미리 짓기 시작. 단기 탄력은 가벼운 이름"],
            ["광", "코히런트 · 루멘텀 · 크레도", "800G에서 1.6T, 3.2T. 중국 규제는 3.2T 공급망을 누가 설계하느냐"],
            ["소재", "두산 전자BG", "광모듈 CCL 3분기 1,059억, 전분기 대비 +115%. 규제 헤드라인과 매출은 다른 층"],
            ["장비", "성호전자", "CPO 정렬. 노출은 맞고, 가격은 111백만주 확인 다음"],
            ["전력", "산일전기", "변압기 중 앞. FY28 16배, 이익 증가가 매출보다 빠름"],
            ["옆길", "Seagate · Western Digital", "금요일 −10%. PER 26배·23배. 메모리와 같은 매수가 아님"],
            ["온도", "오픈AI · 앤트로픽", "토큰이 10배, 15배, 20배로 가는지. 2조 달러는 플랫폼이 돼야 하는 강세"],
        ],
        col_widths=[2.2, 6.6, 8.8],
    )
    n.p("포트의 코너는 이 지도를 따른다. AI 50~60% 안에 삼전·닉스 30% 또는 40%, 소부장 20%. 2차전지 10%, 건설·조선 10%, 현금 20%. 시간이 부족하면 현금 30%. 소부장을 AI 밖에 또 더하면 합이 깨진다.")

    n.h1("메모리만 보면, 피크 다음이다", num="4.")
    n.p("가격의 기본 경로는 지금 100, 6~9개월 뒤 122, 그 1년 뒤 90이다. +22% 뒤에 피크 대비 −26%, 오늘보다는 −10%다. 계약이 가격을 멈추지는 않는다. 2030년 매출의 26%만 미리 정한 밴드 안에 있고, 시장가가 20~30% 빠지면 평균 판매가격은 −18~−27%다.")
    n.table(
        ["", "지금", "피크", "28년 상반기"],
        [
            ["가격", "100", "122", "90"],
            ["분기 영업이익", "65~75조", "100~110조", "60~75조"],
        ],
        col_widths=[4.0, 4.4, 4.4, 4.8],
    )
    n.p("이익은 오를 때 가격보다 크게 오른다. 고정비 때문이다. 내릴 때는 과거처럼 −70~−80%로 무너지는 경로가 기본은 아니다. 다만 Base의 한복판은 피크 대비 약 −36%라, 가격 −26%보다 이익이 더 준다. ASP −25%와 비트 성장 +15~20%를 같이 쓰면 매출은 −14~−10%다. 40조 이하는 가격 급락, 비트 둔화, HBM 가격 하락, 공급 증가, 계약의 방어 약화가 한날에 겹칠 때다.")
    n.callout(
        "주가는 여기서 갈린다",
        [
            "184.2만원은 27년 이익의 4.1배이고, 28년 EPS 28만원으로 두면 이미 6.6배다.",
            "8배면 224만원, 9배면 252만원. 4~6배에 머물면 112~168만원으로 현재보다 낮다.",
            "살 이유는 4.1배가 싸다가 아니다. 28년 분기 이익이 60조 위에 남고, 시장이 8배를 주는가다.",
            "본주와 마이크론의 비교는 27년 4.1배 대 6.0배, 할인 약 31%다. ADR 263만원과 본주 202·219만은 배율 10이 확인된 뒤에만 쓴다.",
        ],
        kind="key",
    )
    n.p("건물 먼저, 장비는 수요 확인 뒤. 마이크론 500억 달러 투자도 그 순서로 읽는다. 상반기 250억 달러에 하반기가 더 크면 연간은 500억 달러 위다. 비트가 내일 그만큼 나오지는 않는다.")

    n.h1("기울기가 바뀌는 자리", num="5.")
    n.table(
        ["보면", "흐름이 가팔라지고", "흐름이 계단으로 남거나 꺾인다"],
        [
            ["금리", "10년물이 5% 아래에 추세적으로 안착", "12월 인상 확률이 다시 산다"],
            ["메모리", "28년 분기 이익이 60조 위", "40조 경로. BNK 2027년 249조는 동종보다 이미 38% 낮다"],
            ["광 · 두산", "1.6T 배치와 CCL 매출이 같이 간다", "규제 헤드라인만으로 CCL을 판다"],
            ["성호전자", "111백만주가 확인되고 CPO 일정이 유지", "7만원, 또는 5~6만원을 30배 하단으로 말한다"],
            ["HDD", "26·27년 계약 물량이 유지", "니어라인 계약이 깨지거나 증설이 앞당겨진다"],
            ["수요 온도", "에이전트 사용이 토큰을 10배에서 15배, 20배로", "앤트로픽 2조를 메모리 기본 이익에 넣는다"],
        ],
        col_widths=[2.8, 7.4, 7.4],
    )
    n.p("앤트로픽의 사모가치 9,650억 달러는 베이스에서 성장률을 21.4%로 올리면 설명된다. 마진 25%와 성장 20%만으로는 약 1.1조 달러다. 2조 달러는 그 위에 플랫폼이 약 1.8배 더 붙는 자리다. 메모리 수요의 온도계로 쓰고, 기본 시나리오의 이익에는 넣지 않는다.")

    n.h1("말로 가져갈 것", num="6.")
    n.callout(
        "네 문장",
        [
            "천장은 10년물 5%대다. 고용 둔화는 장을 받치되, 지수 전체를 한 단 더 올리지 않는다.",
            "돈은 GPU 다음으로 간다. 메모리, 기판, 광, 전력이 한 체인이다.",
            "메모리 가격은 122까지 갔다가 90으로 온다. 이익의 기본은 분기 60~75조이고, 주가 224만원은 그 이익에 8배를 줄 때다.",
            "HDD의 −10%와 두산의 급락과 앤트로픽 2조는 이 체인의 같은 문장이 아니다.",
        ],
        kind="key",
    )

    n.spacer(6)
    n.p(
        "— 10월 2일 큰 흐름. 세부 등식은 검증 노트, 종목별 녹음은 강의노트.",
        size=9.5,
        color=GRAY,
        align="right",
    )
    n.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({OUT_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
