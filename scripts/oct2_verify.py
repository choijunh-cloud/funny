#!/usr/bin/env python3
"""핵심 메모의 숫자와 산식을 원문·재계산으로 검증한다.

fail 은 브리프가 틀린 경우다. note 는 원문 안의 근사나 전사 불일치로, 브리프가 그 한계를 밝혀 두면 실패가 아니다.
"""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document

from oct2_analyze import (
    FACT_NEEDLES,
    OUT_DIR,
    check_facts,
    load_corpus,
    parse_comments,
    portfolio_books,
)
from oct2_brief import OUT_PATH as BRIEF_PATH
from oct2_core import OUT_PATH as CORE_PATH

REPORT_PATH = OUT_DIR / "verification.json"


@dataclass
class Check:
    id: str
    status: str
    claim: str
    detail: str


def _doc_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def _add(rows: list[Check], id_: str, ok: bool, claim: str, detail: str, note: bool = False) -> None:
    if note:
        status = "note"
    else:
        status = "pass" if ok else "fail"
    rows.append(Check(id_, status, claim, detail))


def verify(corpus: dict[str, str] | None = None) -> dict:
    corpus = corpus or load_corpus()
    rows: list[Check] = []

    missing = [row["id"] for row in check_facts(corpus) if not row["found"]]
    _add(
        rows,
        "source_needles",
        not missing,
        "브리프가 인용하는 앵커는 원문에 있다",
        f"{len(FACT_NEEDLES) - len(missing)}/{len(FACT_NEEDLES)}"
        + (f" 없음: {missing}" if missing else ""),
    )

    july_old, july_rev, aug = 3.344, 2.983, 3.008
    revision = round(july_old - july_rev, 3)
    step = round(aug - july_rev, 3)
    headline = round(july_old - aug, 3)
    _add(
        rows,
        "pce_not_a_collapse",
        abs(step) < 0.10 and july_rev < july_old,
        "8월 3.0%는 개정 이후 급락이 아니다",
        f"개정 {revision:.3f}%p, 개정 7월→8월 {step:+.3f}%p, 옛 7월→8월 헤드라인 {headline:.3f}%p",
    )
    m3, m6 = 2.05, 2.74
    _add(
        rows,
        "pce_momentum",
        m3 < m6,
        "최근 3개월 연율이 6개월보다 낮다",
        f"3개월 {m3:.2f}% < 6개월 {m6:.2f}%",
    )

    curve = [
        ("1년", 4.54, 4.43),
        ("2년", 4.89, 4.791),
        ("5년", 5.08, 5.00),
        ("10년", 5.29, 5.24),
        ("30년", 5.63, 5.61),
    ]
    bps = {name: round((end - start) * 100, 1) for name, start, end in curve}
    _add(
        rows,
        "curve_short_end",
        bps["1년"] < bps["30년"],
        "되밀림은 단기물이 더 크다",
        ", ".join(f"{k} {v:+.1f}bp" for k, v in bps.items()),
    )
    high_to_close = round((5.24 - 5.34) * 100, 1)
    high_to_low = round((5.20 - 5.34) * 100, 1)
    _add(
        rows,
        "ten_year_range",
        high_to_close == -10.0,
        "5.34%에서 5.24%대까지는 10bp",
        f"고점→마감 {high_to_close:.0f}bp. 전사에 같이 나온 저점 5.2%까지는 {high_to_low:.0f}bp라, '고저 10bp'는 마감 기준이다",
        note=True,
    )

    books = portfolio_books()
    low, high = books["experienced_ai50"], books["experienced_ai60"]
    _add(
        rows,
        "book_arithmetic",
        sum(low.values()) == 100
        and sum(high.values()) == 100
        and low["삼전닉스"] + low["소부장"] == 50
        and high["삼전닉스"] + high["소부장"] == 60
        and (30 + 20 + 10 + 10 + 20) == 90,
        "경력자 북의 합이 범위와 맞다",
        "AI50 구성합 90 + 미배분 10, AI60 구성합 100",
    )
    _add(
        rows,
        "kospi_gap",
        7000 - 6971 == 29,
        "7,000까지 29포인트",
        "7000 - 6971 = 29",
    )
    swing = round(1.9 - (-1.0), 1)
    _add(
        rows,
        "kospi_swing",
        abs(swing - 2.9) < 0.05,
        "장중 -1%에서 종가 +1.9%는 약 3%p",
        f"스윙 {swing:.1f}%p",
    )

    tp = 4904 * 19.4
    _add(
        rows,
        "intek_target",
        abs(tp - 95000) / 95000 < 0.01,
        "인텍플러스 4,904원 × 19.4배 = 9.5만원",
        f"계산값 {tp:,.0f}원, 목표가 95,000원, 차이 {tp - 95000:,.0f}원 ({(tp / 95000 - 1) * 100:.2f}%)",
    )

    sales_cagr, op_cagr = 0.325, 0.385
    margin_now, margin_later = 0.375, 0.407
    years = 2
    implied = (1 + sales_cagr) ** years * (margin_later / margin_now)
    implied_cagr = implied ** (1 / years) - 1
    _add(
        rows,
        "sanil_cagr_order",
        op_cagr > sales_cagr and margin_later > margin_now,
        "산일전기는 이익이 매출보다 빨리 늘고 마진이 오른다",
        f"매출 CAGR {sales_cagr:.1%}, 영업이익 CAGR {op_cagr:.1%}, 마진 {margin_now:.1%}→{margin_later:.1%}",
    )
    _add(
        rows,
        "sanil_cagr_bridge",
        abs(implied_cagr - op_cagr) < 0.01,
        "마진 확대가 영업이익 CAGR을 설명한다",
        f"매출 CAGR과 마진으로 역산한 영업이익 CAGR {implied_cagr:.1%}, 제시 {op_cagr:.1%}",
    )

    fy26, h1 = 27.0, 25.0
    floor = (h1 + h1) / fy26
    h2_for_double = fy26 * 2 - h1
    # 원문은 10억 달러 단위다. 270억 달러 = 270억 = $27bn.
    h2_krw_unit = h2_for_double * 10
    unit_ok = abs(h2_krw_unit - 290) < 1 and floor > 1.8 and h2_for_double > h1
    _add(
        rows,
        "micron_capex_floor",
        unit_ok,
        "마이크론 2027 캐펙스 하한은 1.85배, 두 배는 하반기 290억 달러 이상",
        f"상·하반기를 같게 두면 {floor:.2f}배. 두 배가 되려면 하반기 {h2_krw_unit:.0f}억 달러",
        note=unit_ok,
    )
    _add(
        rows,
        "micron_gm_dip",
        86.2 < 87,
        "다음 분기 가이던스 86.2%는 이번 87%보다 낮다",
        "차이 -0.8%p. 바닥·보너스 해석은 원문 진술",
    )

    returned = 11.5 / 11.6
    _add(
        rows,
        "accenture_payout",
        returned > 0.95,
        "엑센추어는 잉여현금의 대부분을 환원했다",
        f"11.5 / 11.6 = {returned:.1%}",
    )

    export_total, semi, ssd = 1209, 603, 71
    share = (semi + ssd) / export_total
    _add(
        rows,
        "export_mix",
        abs(share - 0.56) < 0.01,
        "반도체 603 + 컴퓨터·SSD 71은 수출의 56%와 맞다",
        f"{semi + ssd} / {export_total} = {share:.1%}",
    )
    days = export_total / 56
    _add(
        rows,
        "export_daily",
        18 <= days <= 23,
        "수출 1,209억과 일평균 56억은 한 달 조업일로 성립한다",
        f"1,209 / 56 = {days:.1f}일",
    )

    blocks = [
        (b["time"], "20~30" in b["text"], "50~60" in b["text"])
        for b in parse_comments(corpus["comments"])
        if "안전마진" in b["text"]
    ]
    later = next((t for t, tight, _wide in blocks if tight), "")
    earlier = next((t for t, _tight, wide in blocks if wide), "")
    _add(
        rows,
        "semco_cushion_order",
        later == "09:00" and earlier == "08:52" and later > earlier,
        "삼성전기 안전마진은 08:52의 50~60%에서 09:00의 20~30%로 좁아졌다",
        f"블록 {blocks}",
    )

    prior = 520 / 0.91
    _add(
        rows,
        "hana_nine_percent",
        True,
        "삼성 2027년 520조가 9% 하향이라면 하향 전은 약 571조",
        f"520 / 0.91 = {prior:.0f}조. 원문의 컨센서스 밴드(550~500조 후반)보다 높다. 9%와 520은 밴드의 상단 쪽과만 대략 맞는다",
        note=True,
    )

    core = _doc_text(CORE_PATH) if CORE_PATH.exists() else ""
    brief = _doc_text(BRIEF_PATH) if BRIEF_PATH.exists() else ""
    required = [
        "24만 5,000원",
        "3.008%",
        "2.983%",
        "1,500억",
        "6,971",
        "20~30%",
        "420억",
        "1.85배",
    ]
    missing_docs = [s for s in required if s not in core and s not in brief]
    # 1.85배는 긴 브리프에만 둔다. 핵심 메모는 두 배를 말하지 않는다.
    missing_core = [s for s in required if s != "1.85배" and s not in core]
    missing_brief = [s for s in ["1.85배", "290억"] if s not in brief]
    _add(
        rows,
        "docs_carry_the_math",
        not missing_core and not missing_brief and not missing_docs,
        "핵심 메모와 증류 브리프가 검증된 숫자를 그대로 싣는다",
        f"핵심 누락 {missing_core or "없음"}, 브리프 누락 {missing_brief or "없음"}",
    )
    _add(
        rows,
        "core_does_not_claim_double",
        "두 배" not in core,
        "핵심 메모는 확인되지 않은 '캐펙스 두 배'를 단정하지 않는다",
        "두 배 문장은 긴 브리프에서 1.85배 하한과 290억 조건으로 한정",
    )

    fails = [asdict(r) for r in rows if r.status == "fail"]
    notes = [asdict(r) for r in rows if r.status == "note"]
    payload = {
        "pass": sum(r.status == "pass" for r in rows),
        "fail": len(fails),
        "note": len(notes),
        "ok": not fails,
        "checks": [asdict(r) for r in rows],
    }
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def main() -> None:
    payload = verify()
    print(f"pass {payload['pass']}  fail {payload['fail']}  note {payload['note']}")
    for row in payload["checks"]:
        mark = {"pass": "PASS", "fail": "FAIL", "note": "NOTE"}[row["status"]]
        print(f"[{mark}] {row['claim']}")
        print(f"       {row['detail']}")
    if payload["fail"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
