#!/usr/bin/env python3
"""9월 중순 통합브리핑 강의노트(.docx) 생성."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

OUT_PATH = Path("/workspace/lectures/2026년 9월 중순 통합보고서.docx")
CHARTS = Path("/workspace/lectures/assets/sep15")

KR_FONT = "맑은 고딕"
NAVY = RGBColor(0x0F, 0x20, 0x43)
NAVY2 = RGBColor(0x1E, 0x40, 0x7C)
GOLD = RGBColor(0xB8, 0x94, 0x3A)
GRAY = RGBColor(0x4B, 0x55, 0x63)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x16, 0x65, 0x34)
RED = RGBColor(0x99, 0x1B, 0x1B)
AMBER = RGBColor(0x7A, 0x5C, 0x12)
NAVY_HEX, NAVY2_HEX, GOLD_HEX = "0F2043", "1E407C", "B8943A"
LIGHT_HEX, GREEN_HEX, RED_HEX = "EEF2F8", "E8F5E9", "FDECEA"
AMBER_HEX, BLUE_HEX, ROW_HEX, WHITE_HEX = "FFF8E7", "E8F1FB", "F7F9FC", "FFFFFF"


def set_run_font(run, size=11, bold=False, color=DARK, italic=False, font=KR_FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/>'
            f"</w:tcMar>"
        )
    )


def set_table_borders(table, color="D0D7E2", sz="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    tbl_pr.append(
        parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f"</w:tblBorders>"
        )
    )


def set_left_accent(cell, color=NAVY_HEX, sz="24"):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="nil"/><w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="nil"/><w:right w:val="nil"/>'
            f"</w:tcBorders>"
        )
    )


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def cell_text(cell, text, size=10, bold=False, color=DARK, align="left"):
    cell.text = ""
    align_enum = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = align_enum
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1 if i < len(lines) - 1 else 0)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(line)
        set_run_font(run, size=size, bold=bold, color=color)
    set_cell_margins(cell)


class Notes:
    def __init__(self):
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width = Mm(210)
        sec.page_height = Mm(297)
        sec.left_margin = Mm(16)
        sec.right_margin = Mm(16)
        sec.top_margin = Mm(16)
        sec.bottom_margin = Mm(16)
        normal = self.doc.styles["Normal"]
        normal.font.name = KR_FONT
        normal.font.size = Pt(11)
        normal.font.color.rgb = DARK
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.18
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("2026년 9월 중순 통합보고서  ·  AI 속도조절 · Token · 매크로 · 메모리")
        set_run_font(r, size=8.5, color=GRAY)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("원문 17개 + Quick 코멘트 통합  ·  숫자·시나리오는 원문 기준  ·  ")
        set_run_font(r, size=8, color=GRAY)
        fp._p.append(
            parse_xml(
                f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE ">'
                f'<w:r><w:rPr><w:sz w:val="16"/><w:color w:val="4B5563"/>'
                f'<w:rFonts w:ascii="{KR_FONT}" w:hAnsi="{KR_FONT}" w:eastAsia="{KR_FONT}"/></w:rPr>'
                f"<w:t></w:t></w:r></w:fldSimple>"
            )
        )
        core = self.doc.core_properties
        core.title = "2026년 9월 중순 통합보고서"
        core.author = "준혁"
        core.subject = "AI 속도조절 · Token Economy · 매크로 New Normal · 메모리"

    def p(self, text, size=11, bold=False, color=DARK, space_after=6, space_before=0, align="left"):
        para = self.doc.add_paragraph()
        para.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(space_before)
        para.paragraph_format.line_spacing = 1.18
        run = para.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
        return para

    def h1(self, text, num=None):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(14)
        para.paragraph_format.space_after = Pt(8)
        if num:
            run = para.add_run(f"{num}  ")
            set_run_font(run, size=16, bold=True, color=GOLD)
        run = para.add_run(text)
        set_run_font(run, size=16, bold=True, color=NAVY)
        para._p.get_or_add_pPr().append(
            parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="12" w:space="4" w:color="{NAVY_HEX}"/>'
                f"</w:pBdr>"
            )
        )
        return para

    def h2(self, text):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(4)
        run = para.add_run(text)
        set_run_font(run, size=13, bold=True, color=NAVY2)
        return para

    def bullet(self, text, level=0, size=11):
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.55 + level * 0.45)
        para.paragraph_format.first_line_indent = Cm(-0.35)
        para.paragraph_format.space_after = Pt(2.5)
        para.paragraph_format.line_spacing = 1.15
        mark = "• " if level == 0 else "– "
        run = para.add_run(mark)
        set_run_font(run, size=size, color=NAVY2 if level == 0 else GRAY)
        run = para.add_run(text)
        set_run_font(run, size=size, color=DARK)
        return para

    def flow(self, items, size=10.5):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(8)
        for i, item in enumerate(items):
            if i:
                run = para.add_run("  →  ")
                set_run_font(run, size=size, bold=True, color=GOLD)
            run = para.add_run(item)
            set_run_font(run, size=size, bold=True, color=NAVY)
        return para

    def callout(self, title, body, kind="key"):
        palette = {
            "key": (NAVY_HEX, LIGHT_HEX, NAVY),
            "bull": ("166534", GREEN_HEX, GREEN),
            "bear": ("991B1B", RED_HEX, RED),
            "note": (GOLD_HEX, AMBER_HEX, AMBER),
            "blue": (NAVY2_HEX, BLUE_HEX, NAVY2),
        }
        accent, fill, title_color = palette[kind]
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        shade_cell(cell, fill)
        set_left_accent(cell, accent, sz="28")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r = p1.add_run(title)
        set_run_font(r, size=10, bold=True, color=title_color)
        if isinstance(body, str):
            body = [body]
        for line in body:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(line)
            set_run_font(r, size=10.5, color=DARK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def table(self, headers, rows, col_widths=None, first_col_bold=True):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="D5DCE6", sz="4")
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            shade_cell(cell, NAVY_HEX)
            cell_text(cell, h, size=9, bold=True, color=WHITE, align="center")
        prevent_row_split(table.rows[0])
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = table.rows[r_i + 1].cells[c_i]
                shade_cell(cell, ROW_HEX if r_i % 2 else WHITE_HEX)
                align = "left" if c_i == 0 else "center"
                cell_text(cell, str(val), size=9, bold=first_col_bold and c_i == 0, color=DARK, align=align)
            prevent_row_split(table.rows[r_i + 1])
        if col_widths:
            for row in table.rows:
                for i, w in enumerate(col_widths):
                    row.cells[i].width = Cm(w)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)
        return table

    def image(self, name, width=16.2, caption=None):
        path = CHARTS / name
        if not path.exists():
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Cm(width))
        if caption:
            self.p(caption, size=8.5, color=GRAY, space_after=10, align="center")

    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))


def build():
    n = Notes()
    n.p("CONFIDENTIAL  ·  내부용 통합보고서", size=10, color=GOLD, align="center", space_after=18)
    n.p("2026년 9월 중순", size=13, bold=True, color=GOLD, align="center", space_after=4)
    n.p("통합보고서", size=28, bold=True, color=NAVY, align="center", space_after=6)
    n.p("AI 속도조절론 · Token Economy · 매크로 New Normal · 메모리", size=14, bold=True, color=NAVY2, align="center", space_after=8)
    n.p("원문 17개와 Quick 코멘트를 하나의 문서로 통합", size=11, color=GRAY, align="center", space_after=4)
    n.p("기준일  2026. 9. 11–15", size=10.5, color=GRAY, align="center", space_after=14)

    n.callout(
        "이 보고서의 결론",
        [
            "「Pacing the Frontier」는 AI 인프라 Peak-out이 아니다. 더 강한 모델에서 더 많은 토큰으로 투자 축이 이동한다.",
            "9/14 SOX 급락은 수요 훼손보다 속도조절 우려 + 10년물 5% + 유가 $100대가 겹친 Risk-off다.",
            "지금은 Peak-out이 아니라 Timing Risk를 점검할 단계다. KEY는 Frontier가 늦어도 현재 AI+Agent만으로 Capex가 지속되느냐다.",
            "이 파일 하나만 보시면 됩니다. 차트·도표·종목 논점은 본문에 모두 들어 있습니다.",
        ],
    )

    n.h2("목차")
    n.table(
        ["파트", "주제", "가져갈 한 줄"],
        [
            ["0", "숫자 한 장", "7/31 이후 금리·유가 악재에도 메모리가 이김"],
            ["1", "속도조절·종말론", "출시 속도 ≠ 인프라 속도. 킬 스위치 6층"],
            ["2", "Token / NVDA vs Memory", "논리 70~80% 타당. Inference 점유율이 진짜 위험"],
            ["3", "KEY 질문", "생산성 ROI · RSI · 가드레일=추가 연산"],
            ["4", "매크로", "바닥↑와 장기채 가격↑는 모순 아님. 유가가 Key"],
            ["5", "메모리 밸류·가격", "본주 할인 / ADR 40% 프리미엄 too high"],
            ["6", "전력", "상품은 효율이 아니라 Time-to-Power"],
            ["7", "종목", "SFA · PSK · BH · 휴머노이드 · Oracle · 무라타"],
            ["8", "체크리스트", "가이던스→발주→가동→출시 연속 지연만 Downside"],
            ["부록", "자료 목록", "원문 17개와 이 보고서 파트의 대응"],
        ],
        col_widths=[2.0, 5.2, 10.4],
    )
    n.page_break()

    # 0
    n.h1("오늘 숫자 한 장", num="0.")
    n.image("01_july_sep_returns.png", caption="7/31→9/1. 유가 +18.8% · 2년물 +7.9% · 10년물 +4.8%에도 샌디스크 +34.5% · 하이닉스 ADR +32.2%.")
    n.image("02_macro_shock.png")
    n.table(
        ["항목", "9/11", "9/14", "읽는 법"],
        [
            ["Sandisk FY27 PER", "8.1배 ($1,633)", "7.7배 ($1,552)", "EPS $201"],
            ["Micron CY27 PER", "6.5배 ($975)", "6.2배 ($924)", "EPS $150"],
            ["하이닉스 ADR 27E", "5.8~6.5배 / $190", "5.4~6.0배 / $176", "과거 vs MU −20~−50%"],
            ["하이닉스 본주", "181.2만 / 27E 4.6배*", "169.7만 / 4.3배*", "ADR 프리미엄 ~40% too high"],
            ["삼성전자", "25.95만 / 4배대", "24.9만 / 4.2배", "MU 대비 ~30%대 할인"],
            ["S&P Fwd PER", "26E 21 / 27E 18.5", "—", "EY 21배=4.76%, 10Y 5% 대비 −0.24%p"],
        ],
        col_widths=[3.6, 4.4, 4.6, 5.0],
    )
    n.p("*본주 27E는 환율 하향 반영. 원화 강세 10% ≈ 삼전닉스 −12%.", size=9.5, color=GRAY)

    # 1
    n.h1("속도조절론 · 종말론", num="1.")
    n.callout(
        "Amodei가 말한 것 / 말하지 않은 것",
        [
            "말한 것: 위험이 있다. 속도가 예상보다 빨랐다. 모델마다 안전 테스트. 킬 스위치는 좋은 아이디어일 수 있다. 독립 평가자=식품 검사관. 중국이 안 멈추면 미국만 멈추기 어렵다.",
            "말하지 않은 것: 개발 전면 중단, GPU·DC 축소. Anthropic은 최대 $517B 컴퓨팅 계약·14.8GW, Google Cloud 5년 $200B. Pacing = 모델 출시 속도.",
        ],
        kind="key",
    )
    n.flow(["챗봇=답하는 도구", "Agent=행동하는 주체", "검색·코드·CRM·금융", "위험의 차원이 달라짐"])
    n.h2("킬 스위치 6층")
    n.table(
        ["층", "대상", "의미"],
        [
            ["①", "모델", "위험한 행동 차단"],
            ["②", "Agent", "특정 행동 권한 제한"],
            ["③", "시스템", "네트워크·컴퓨터 접근 차단"],
            ["④", "인프라", "GPU/서버 실행 중단"],
            ["⑤", "인간", "최종 승인권"],
            ["⑥", "제3자", "독립 안전성 검사 (특히 강조)"],
        ],
        col_widths=[2.2, 4.0, 11.4],
    )
    n.h2("종말론 세 겹 + 네 독해")
    n.bullet("진짜 우려 + 전략적 프레이밍 + PR 과장이 섞여 있다 — 대중의 1차 반응.")
    n.table(
        ["프레임", "한 줄"],
        [
            ["인류애 / 당위", "1953 Atoms for Peace ↔ 2026 AI for Humanity. 능력이 통제를 앞지름."],
            ["실패 대비 서사", "LLM을 AGI에 연결. $2T+가 AGI 전제. 실패 시 안전/가격 양다리."],
            ["순리론", "위험은 진심, 가드레일도 옳다. 그래도 Picks and Shovels."],
            ["냉소", "시진핑 BS → 미국이 진지하게 받아들일 이유 소멸. 동시 감속은 불안정."],
        ],
        col_widths=[4.2, 13.4],
    )
    n.h2("두 견해를 대치")
    n.p("견해 A — 단기 Timing 지연: 투자 종료가 아니라 집행 시점 불확실. 수요의 방향이 아니라 시간표.")
    n.p("견해 B — RSI + 빅테크 우위: 병목이 인력→컴퓨팅. 규제 비용은 수직계열화 가치. 모델 기업은 클라우드·고객 플랫폼 필요.")
    n.h2("포지션 맵")
    n.table(
        ["주체", "입장", "함의"],
        [
            ["Trump", "Whoever wins AI, wins. 늦출 이유 없다", "투자 지속. 주말 Perp 매도 완충"],
            ["Amodei", "안전장치가 속도를 못 따라감", "규제 설계권 · 공동 pacing"],
            ["Altman·Musk", "독립 평가 동조. 동시에 Grok 5=AGI", "가드레일과 경쟁이 동시에"],
            ["백악관", "자율규제·시장규율", "강제보다 성장 우선"],
            ["의회", "민주=감독 / 공화=혁신", "규제 불확실성 지속"],
            ["Karp", "적대국이 안 멈추면 미국만 감속은 안보 위험", "단독 감속 불가"],
        ],
        col_widths=[3.2, 7.6, 6.8],
    )
    n.callout(
        "정책 한 줄",
        ["안전장치는 보완하되 패권 경쟁과 투자는 늦추지 않는다. 주말 뉴스의 펀더멘털 충격은 제한적일 가능성."],
        kind="blue",
    )
    n.image("11_perp_meaning.png", caption="SKHYNIX Perp −4.35%는 현물 예고가 아님. 야간 심리 온도계. 7/28 Oracle 왜곡으로 18~20% 순간 폭락·$57M 청산.")

    # 2
    n.h1("Token Economy · NVIDIA vs Memory", num="2.")
    n.flow(["Frontier 속도조절", "최신 GPU 긴급도 ↓", "Rubin 채택·마진 불확실"])
    n.flow(["추론비용↓ + Agent", "Token 폭증", "KV Cache / Context", "Server DRAM + NAND"])
    n.image("12_capex_formula.png", caption="AI Capex ≈ Workload × Tokens/Workload × Compute/Token × $/Compute. 단가↓여도 총수요↑ 가능.")
    n.callout(
        "가장 강한 논리",
        ["모델 성능 증가율 ↓ ≠ AI 사용량 증가율 ↓. 추론 비용이 내려가면 가격탄력성이 작동할 수 있다."],
        kind="bull",
    )
    n.callout(
        "가장 큰 약점",
        [
            "원문은 NVIDIA를 Frontier GPU 회사로 너무 강하게 본다.",
            "진짜 위험은 Inference 감소가 아니라 증가분 중 NVIDIA 점유율 하락.",
            "Yes: reasoning / long-context / multimodal / agentic은 고연산. No: 단순 추론은 ASIC·저가 GPU·CPU.",
        ],
        kind="bear",
    )
    n.callout(
        "Memory 논리는 오히려 강하다",
        [
            "AI=HBM이면 Frontier↓→NVDA↓→HBM↓→MU↓.",
            "AI=HBM+Commodity DRAM+NAND+eSSD+CXL이면 Token→Context→KV→Server DRAM→Storage.",
            "서버가 GPU+HBM+CPU+DDR5/LPDDR+CXL+SSD 계층이 되면 TAM 확장. DRAM도 AI.",
        ],
        kind="note",
    )
    n.h2("ASIC은 GPU를 죽이나")
    n.p("Agent 상시가동 → 추론이 원가의 핵심 → 패턴이 안정되면 TPU/Trainium/Inferentia/Maia 유인. 표준화 전·복잡한 초기 업무는 범용 GPU. 대규모·반복 추론부터 비중 상승. HBM 외 DRAM·네트워크·패키징 장기화.")
    n.h2("젠슨 황 (9/11 GS) 압축")
    n.table(
        ["#", "포인트"],
        [
            ["1", "2030 인프라 $3~4T. Retrieval→Generative. Extreme co-design"],
            ["2", "다음 ROI 대형 시장 = 사이버보안"],
            ["3", "칩 회사가 아니라 AI 팩토리. NVL72 월 +27%"],
            ["4", "폐쇄·오픈 모델·퀀트·제약·네오클라우드 모두 수혜"],
            ["5", "패키징·DRAM·LPDDR·커넥터·웨이퍼 tight. 더 큰 병목은 land/power/shell"],
            ["6–7", "네오클라우드 backlog 수천억$. 내년 매출 +70%, 무제약 수요 +100%+"],
            ["8–9", "2GW=$80B. AI native 자금 2/3가 컴퓨팅. 1GW 시스템 ~$60B"],
            ["10", "Physical AI: 자율주행→물류→조작 로봇(~2년)→6G(~5년)"],
        ],
        col_widths=[2.4, 15.2],
    )

    # 3
    n.h1("KEY — 현재 속도로도 생산성은 충분한가", num="3.")
    n.callout(
        "검증이 필요한 문장",
        [
            "Frontier가 늦어져도 AI Capex는 계속되는가.",
            "AGI급 속도전이 없어도, 지금 수준 + Agentic End-to-end만으로 기업 이득이 투자를 지속시킬 것인가.",
            "RSI 초기(AI가 코드·실험·분석 등 R&D에 참여)만으로도 컴퓨팅 선점은 원래 속도대로 늘어야 한다.",
        ],
        kind="note",
    )
    n.flow(["안전성 강화", "검증·모니터링", "Agent 전후 테스트", "상시 감시", "연산 ↑"])
    n.h2("같은 주, 반대처럼 보이는 뉴스")
    n.table(
        ["뉴스", "읽기"],
        [
            ["Grok 4.9=Astra/Fable급, 5=AGI (probably/maybe)", "목표이지 벤치마크 확정 아님. 경쟁은 안 끝남"],
            ["OpenAI Glass Imaging $3억+", "Vision→Agent. Physical interface"],
            ["Apple Siri AI", "기기가 Agent 플랫폼"],
            ["Google 뉴멕시코 DC (Permian)", "병목이 GPU에서 전력·부지"],
            ["BofA 등 $3T+ 속 노이즈", "가이던스 확인 전이면 저가매수 시각도 존재"],
        ],
        col_widths=[7.4, 10.2],
    )

    # 4
    n.h1("매크로 New Normal", num="4.")
    n.p("9월 인상 자체는 상당 반영. 핵심은 속도·Ceiling, 유가→기대인플레→전가→Core CPI.")
    n.table(
        ["상황", "10년물", "장기채 가격"],
        [
            ["현재", "5.0%", "100"],
            ["신뢰받는 긴축 → 물가 기대↓", "4.5%↓", "상승"],
            ["경기침체까지", "4.0%↓", "더 상승"],
            ["유가+재정+국채공급 → 기간 프리미엄↑", "금리↑", "하락"],
        ],
        col_widths=[8.0, 4.4, 5.2],
    )
    n.callout(
        "모순이 아닌 이유",
        ["바닥이 높아졌다 = 1~2%로 돌아가기 어렵다. 그래도 5%→4%면 가격은 오른다. 과거 1%→0.5% 랠리의 축소판."],
        kind="blue",
    )
    n.p("30년물 5.3%대에서 연기금·보험·해외가 받음: 낙찰 5.308%, 응찰 2.61배, 간접 79.5%, PD 2.2%, WI −2.7bp. 다음 시험대는 9/24 바이백(잔존 20~30년).")
    n.image("06_per_history.png", caption="현재 PER 20배대 + 10년물 4.8~5.0%는 역사적으로 이례적.")
    n.image("05_per_vs_yield.png", caption="S&P 26E 컨센 EPS ~$365, 27E $415~420 (+14~15%). GS는 $340/$385로 더 보수.")
    n.h2("유가 경로")
    n.table(
        ["경로", "금리", "장기금리"],
        [
            ["유가 안정 + 기대 안정", "9월 후 추가 인상 필요↓", "안정"],
            ["유가 고공 + 기대 고착", "10월 이후 추가↑ 가능", "재상승"],
        ],
        col_widths=[6.4, 6.0, 5.2],
    )
    n.p("9/11: CPI 부합, Core 상회. 그런데 SOX +1.8%. 금리 경로 스케치: 12월→9월 인상, 10월 연속 가능성 낮음. 채권은 단기↑↔장기↓가 동시에 가능한 구간.")
    n.h2("엔화")
    n.bullet("BoJ 인상은 반복 반영. 중기 추세 전환으로 보기 어려움.")
    n.bullet("엔캐리 붕괴의 역사적 트리거는 정책이 아니라 큰 해외 손실.")
    n.bullet("지금은 숏 과다의 되돌림(2024.8 오해와 유사). 원/달러 추가 하락 재료로 보기 어려움.")
    n.h2("9/14 미국장")
    n.flow(["속도조절", "10년물 5%", "유가 $100~110", "SOX 차익실현"])
    n.p("낙폭은 컸으나 7/29 바닥은 멀고, 급반등 대비 소폭 하락 반전. MS +2% · GOOGL +3.2% · IGV +5%. 곡괭이를 좀 줄이는 심리는 타당, 그 이상 액션의 증거는 대기.")

    # 5
    n.h1("메모리 밸류 · 가격 · 리드타임", num="5.")
    n.image("03_memory_per.png")
    n.table(
        ["", "26년 OP / EPS", "27년 OP / EPS"],
        [
            ["SK하이닉스", "266조 / 350K", "392조 / 436K → 392K? (환율 −10% 가능)"],
            ["삼성전자", "392조 / 48.1K", "543조 / 66.4K → 60K?"],
        ],
        col_widths=[3.6, 6.6, 7.4],
    )
    n.callout(
        "ADR vs 본주",
        [
            "ADR 본주 대비 ~40% 프리미엄 = too high.",
            "9/14 기준 30%면 본주 182만, 20%면 198만(희소성상 단기는 어려움).",
            "보수: 27년 성장 0, 26년 PER 6~7배 → 하이닉스 210~245만, 삼성 28.9~33.7만.",
        ],
        kind="note",
    )
    n.image("07_apple_memory.png", caption="1Q27 DRAM ~$2.0/Gb, NAND ~$0.33/Gb (+30~40%). iPhone 18 Pro 256GB 메모리 원가 YoY +400%, BOM 10%→34%.")
    n.bullet("서버·HBM capa 우선, 고객 확보 물량 신청의 60~70%. NAND 웨이퍼 계획 축소(삼성 490→468만장 등).")
    n.bullet("폴더블 Duo $1,999 = 희소 메모리를 초고가에 배분. 단순 폴더블이면 Bear, AI×대화면이면 Upgrade Cycle.")
    n.image("04_lead_times.png", caption="GPU만 균형. ABF 48–56주, HDD 50주, MLCC 30주, DRAM 20주, NAND 16주.")
    n.p("무라타: AI MLCC 수요 최소 2028. 2년 800억엔 + 이후 3~5년 추가 건물 검토. 가격은 장기 신뢰(잦은 인상=신규 진입). 하이엔드 부하로 로우엔드 점유 우선순위↓. 영향은 28년 가시성 확인, 경쟁사 가격만 경계 — 중립~긍정.")
    n.h2("온디바이스 — NPU만 보면 절반")
    n.table(
        ["주장", "판단"],
        [
            ["PC에서 작은 LLM은 처참하다", "대체로 맞음"],
            ["그러므로 온디바이스는 쓸모없다", "틀림"],
            ["칩에 최적화하면 쓸 만하다", "맞음"],
            ["아이폰이 ChatGPT를 대체한다", "아님"],
            ["복잡한 AI는 클라우드와 결합", "맞음"],
        ],
        col_widths=[9.0, 8.6],
    )
    n.p("3B + 2-bit + LoRA + Neural Engine. AFM 3는 모델을 NAND에 두고 선택 활성화. 기능↑ → NAND↑. LPDDR + NAND + Memory Tiering.")

    # 6
    n.h1("전력 · 연료전지", num="6.")
    n.p("연료전지의 상품은 효율이 아니라 Time-to-Power. 원전은 규모·24시간·저탄소. LNG가 가장 현실적 경쟁자. PJM 혼잡비용 26H1 약 $60억. Google 615MW 25년 구매.")
    n.table(
        ["규모", "해법"],
        [
            ["50~200MW", "연료전지 매우 매력 (PAFC/SOFC)"],
            ["200~500MW", "FC 100~300MW + LNG/Grid"],
            ["500MW~1GW", "원전/LNG/대형 Grid 중심, FC는 보완"],
            ["1GW+", "원전+Grid+LNG/재생+ESS. FC만으로는 비경제"],
        ],
        col_widths=[4.4, 13.2],
    )
    n.image("09_doosan_tp.png", caption="국내 연료전지 → 미국 AI 전력 인프라로의 재평가. 7~8만은 추가 수주·잔고·가동률·OPM이 확인돼야 함.")
    n.bullet("미국 DC ~5,000억, 내년 3월~, OPM 6~7%, PAFC CAPA 350MW. 데이터센터 규제=금지가 아니라 요금 부담의 정치적 배분.")
    n.bullet("웨스팅하우스 15% 타진 = 원전 축. 연료전지와 시간축이 다른 보완재.")

    # 7
    n.h1("종목 · 밸류체인", num="7.")
    n.h2("SFA반도체")
    n.flow(["삼성 DRAM", "온양 WT", "천안·온양 패키징", "Final Test", "HBM"])
    n.flow(["온양 HBM 전환", "DDR5 외주↑", "SFA 한국 패키징", "필리핀 Final Test"])
    n.bullet("베트남 증설 우려는 LPDDR·UFS·MCP 내재화 테스트 중심 가능성. 필리핀은 재인증+가동률 개런티.")
    n.bullet("삼성 외주정책 사이클 주식. 27/28 추정을 믿으면 +30%에 PER 20배 전후 → 9~12만원.")
    n.h2("PSK / PSK홀딩스")
    n.image("08_psk_growth.png", caption="26E/27E 매출 2,792억(+34%) / 4,014억(+44%), OP 1,088(+48%) / 1,648(+52%), OPM 39→41%. 컨센 TP 19.5만 ≈ 27 PER 15배.")
    n.table(
        ["장비", "포인트"],
        [
            ["Dry Strip SUPRA", "PR 플라즈마 제거. 20년+ 점유 1위. HBM·GAA에서 Wet보다 손상↓"],
            ["Descum ECOLITE", "수 nm 잔류막. RDL/Bump 수율. HBM·FO·Chiplet·PLP"],
            ["Fluxless GENEVA", "미세 Pitch에서 Flux 잔류=수율 킬러. Warpage·Throughput"],
        ],
        col_widths=[4.6, 13.0],
    )
    n.p("3Q에 2Q 이연 120억. 26H2~27 발주 재개(OSAT, TSMC, M15X·P4). 신규축 Intel EMIB-T 27H2.")
    n.h2("비에이치")
    n.bullet("듀오 면적↑ → RFPCB 면적 2.8배, ASP 2배. 대신 TP 4.1만=27E×14.5배. 타사 컨센 2만 중반.")
    n.bullet("3만원대에서 안전마진 20%(매크로 우려 시 30%). 폴더블 28년 1,500만대≈아이폰 7% → 매출 ~20%, 이익 10%대. 로봇은 숫자 전엔 배수 요인. 넘버가 더 뚜렷한 대안도 점검.")
    n.h2("휴머노이드")
    n.table(
        ["이름", "포인트"],
        [
            ["삼현", "SEC 액추에이터 인증 진행. Figure AI 4세대(2027) 모터. CES 2027"],
            ["레인보우", "미국 생산시설, 수출 재개 보도"],
            ["로보티즈 · SPG", "테마 대표 / 감속기"],
            ["SDS RX", "계열 공장 실증이 상용화 속도의 무기"],
        ],
        col_widths=[3.8, 13.8],
    )
    n.h2("Oracle")
    n.image("10_oracle.png", caption="RPO $664B, 신규 $30B, 50%가 36개월 내 전환. FY27 Capex $90~95B 중 순현금 $70B 이하(선급금·BYOH $20~25B).")
    n.h2("곁가지")
    n.bullet("사모대출: 소프트웨어 대체 우려 후퇴=심리. 현금흐름·매각·차환으로 확인. 인프라 금융의 낙관이 금리를 보상하는지도 점검.")
    n.bullet("신정법: 산술 조 단위 가능. 동양생명 1,400→70억. 불확실성 프리미엄. 더 깊은 이슈는 은행 AI 데이터 활용. 금융주 급락 시 이 뉴스 vs 금리·유가 디레이팅을 분리.")

    # 8
    n.h1("체크리스트", num="8.")
    n.callout(
        "네 가지가 유지되면 속도조절=HBM Peak-out은 약해진다",
        [
            "① 하이퍼스케일러 CapEx 가이던스   ② GPU/HBM 발주·선급금",
            "③ DC 착공·가동 일정   ④ 모델 출시와 안전 검증 기간",
            "가이던스→발주→가동에서 연속 지연이 나오면 그때는 Timing Downside.",
        ],
        kind="key",
    )
    n.table(
        ["우선", "항목", "왜"],
        [
            ["1", "빅테크 CapEx 가이던스", "펀더멘털 vs 멀티플의 판정자"],
            ["2", "실제 발주 / 선급금", "가이던스와 집행의 갭"],
            ["3", "10년물 vs 5.3% 수요대, 유가 $100", "성장주 할인율"],
            ["4", "HBM4 주문/가격, 모바일 vs 서버", "업사이클 지속성"],
            ["5", "강제 테스트·독립감사 법제화", "규제 논의가 ①→⑤로 실체화하나"],
            ["6", "기업 AI 사용·생산성 케이스", "KEY 질문의 유일한 실증"],
        ],
        col_widths=[2.0, 6.8, 8.8],
    )
    n.callout(
        "닫는 문장",
        [
            "지금은 Peak-out을 논할 단계가 아니라 Timing Risk를 점검할 단계.",
            "AI 성장률보다 중요한 것은 CapEx 집행 속도. 종말론보다 중요한 것은 생산성의 확장.",
            "그 확장이 숫자로 나타나는지 — 가이던스, 발주, 토큰, 메모리 가격 — 를 보면 된다.",
        ],
        kind="bull",
    )
    n.page_break()
    n.h1("부록 — 원문과 이 보고서의 대응", num="A.")
    n.p("아래 원문을 이 한 권에 통합했습니다. 숫자와 시나리오는 원문 표기를 재구성한 것입니다.")
    n.table(
        ["원문", "이 보고서"],
        [
            ["AI 속도조절론의 2가지 다른 견해", "1. 견해 A(Timing) vs 견해 B(RSI·빅테크)"],
            ["AI종말론, 속도조절론, 생산성 이득", "3. KEY 질문. 표지 결론문"],
            ["Pacing the Frontier", "2. Token Economy · NVDA vs Memory"],
            ["AI 종말론까지 말이 나오면", "1. 네 독해 프레임"],
            ["TalkFile_ai개발속도 조절론", "1. Perp, Trump, 포지션 맵"],
            ["9월 14일 미국장 & 생산성의 확장", "3–4. 9/14 Risk-off, KEY"],
            ["Where is market headed / 매크로 9.14", "4. 금리 New Normal, 유가, 엔화"],
            ["9월 13일 장기금리 · Astra·ASIC · 젠슨", "2. ASIC, 젠슨 10포인트 / 4. 장기채"],
            ["9월 11·12일 강의", "4–5. CPI, SOX, 애플, Oracle"],
            ["7월 말 대비 달라진 것", "0. 수익률 차트"],
            ["S&P500 PER vs 10년물", "4. PER 역사 · EY vs 5%"],
            ["메모리 가격 밸류에이션 9/11·9/14", "5. PER 비교, ADR 프리미엄"],
            ["연료전지", "6. Time-to-Power, 두산"],
            ["SFA반도체", "7. DDR5 외주 · 온양 HBM"],
            ["Quick 코멘트 (PSK, BH, 휴머노이드, 무라타,\n신정법, 사모대출, Grok, Apple 가격 등)", "5·7·곁가지"],
        ],
        col_widths=[8.6, 9.0],
    )
    n.callout(
        "이 파일만 보시면 됩니다",
        [
            "차트 12장, 표, 흐름도, 종목 논점을 이 워드 한 권에 넣었습니다.",
            "재생성: python3 /workspace/scripts/sep15_charts.py && python3 /workspace/scripts/generate_sep15_briefing.py",
        ],
        kind="key",
    )
    n.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
